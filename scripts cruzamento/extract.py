"""Le as bases de faturamento e agrega, por chave, a ultima venda (data, vendedores do dia, cliente, cnpj).

Chaves:
  root  -> raiz do CNPJ (8 digitos)
  cli   -> (CLIENTE, UF)
  desc  -> (DES_PROD, UF)
Cada valor: [data_int, {vendedor: valor_no_dia}, cliente, cnpj, segmento]
"""
import os, sys, glob, pickle, time, datetime, re
from multiprocessing import Pool
from common import BASE, AGG

# todas as bases "fat *.xlsx" da pasta Bases Claude (bases novas entram automaticamente)
FILES = sorted(os.path.basename(p) for p in glob.glob(os.path.join(BASE, "fat *.xlsx")))
NEED = ["EMISSAO", "NOME_VENDEDOR", "CNPJ", "Estado", "CLIENTE", "DES_PROD", "VLR_TOTAL", "SEGMENTO"]
NONDIG = re.compile(r"\D")


def to_int_date(v):
    if v is None:
        return None
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.year * 10000 + v.month * 100 + v.day
    s = str(v).strip()
    if len(s) >= 10 and s[2] == "/" and s[5] == "/":
        try:
            return int(s[6:10]) * 10000 + int(s[3:5]) * 100 + int(s[0:2])
        except ValueError:
            return None
    return None


def upd(d, key, dt, vend, val, cli, cnpj, seg):
    cur = d.get(key)
    if cur is None or dt > cur[0]:
        d[key] = [dt, {vend: val}, cli, cnpj, seg]
    elif dt == cur[0]:
        cur[1][vend] = cur[1].get(vend, 0.0) + val


def work(job):
    fname, sheet = job
    import openpyxl
    t = time.time()
    wb = openpyxl.load_workbook(os.path.join(BASE, fname), read_only=True)
    ws = wb[sheet]
    it = ws.iter_rows(values_only=True)
    hdr = [str(h).strip() if h is not None else "" for h in next(it, ())]
    missing = [n for n in NEED if n not in hdr]
    if missing:
        wb.close()
        return f"{fname} [{sheet}] IGNORADA - colunas nao encontradas no cabecalho: {missing}"
    idx = {}
    for n in NEED:
        idx[n] = hdr.index(n)  # primeira ocorrencia (2026 tem 'MUNICIPIO' extra, nao usado)
    iE, iV, iC, iU, iCl, iP, iT, iS = (idx[n] for n in NEED)
    root, cli, desc = {}, {}, {}
    n = bad = 0
    for r in it:
        n += 1
        dt = to_int_date(r[iE])
        if dt is None:
            bad += 1
            continue
        vend = (r[iV] or "").strip() if isinstance(r[iV], str) else ("" if r[iV] is None else str(r[iV]))
        cnpj = NONDIG.sub("", str(r[iC] or ""))
        uf = (r[iU] or "").strip() if isinstance(r[iU], str) else ""
        cl = (r[iCl] or "").strip() if isinstance(r[iCl], str) else str(r[iCl] or "")
        pr = (r[iP] or "").strip() if isinstance(r[iP], str) else str(r[iP] or "")
        seg = (r[iS] or "").strip() if isinstance(r[iS], str) else ""
        try:
            val = float(r[iT] or 0)
        except (TypeError, ValueError):
            val = 0.0
        if len(cnpj) == 14:
            upd(root, cnpj[:8], dt, vend, val, cl, cnpj, seg)
        if cl:
            upd(cli, (cl, uf), dt, vend, val, cl, cnpj, seg)
        if pr:
            upd(desc, (pr, uf), dt, vend, val, cl, cnpj, seg)
    wb.close()
    out = os.path.join(AGG, f"{fname}__{sheet}.pkl")
    with open(out, "wb") as f:
        pickle.dump({"root": root, "cli": cli, "desc": desc, "rows": n, "bad": bad}, f, protocol=pickle.HIGHEST_PROTOCOL)
    return f"{fname} [{sheet}] linhas={n} sem_data={bad} raizes={len(root)} clientes={len(cli)} descs={len(desc)} t={time.time()-t:.0f}s"


def jobs():
    import openpyxl
    js = []
    for f in FILES:
        wb = openpyxl.load_workbook(os.path.join(BASE, f), read_only=True)
        for s in wb.sheetnames:
            js.append((f, s))
        wb.close()
    # maiores primeiro
    js.sort(key=lambda j: -os.path.getsize(os.path.join(BASE, j[0])))
    return js


if __name__ == "__main__":
    os.makedirs(AGG, exist_ok=True)
    for old in glob.glob(os.path.join(AGG, "*.pkl")):  # limpa a rodada anterior
        os.remove(old)
    if not FILES:
        sys.exit(f"Nenhuma base 'fat *.xlsx' encontrada em {BASE}")
    js = jobs()
    print("abas:", js, flush=True)
    t = time.time()
    with Pool(int(sys.argv[1]) if len(sys.argv) > 1 else 4) as p:
        for msg in p.imap_unordered(work, js):
            print(msg, flush=True)
    print(f"TOTAL {time.time()-t:.0f}s", flush=True)
