import openpyxl, zipfile, re, sys, datetime as dt
a,b=sys.argv[1],sys.argv[2]
RES={'RESUMO SUPERMERCADO (REUNIAO)','"FARMACIA CONSOLIDADO RESUMO"','"FARMACIA VF RESUMO"'}
for f in (a,b):
    z=zipfile.ZipFile(f); n=z.namelist()
    pts=[x for x in n if x.startswith('xl/pivotTables/pivotTable') and x.endswith('.xml')]
    names=sorted(re.search(r' name="([^"]+)"',z.read(x).decode('utf8')).group(1) for x in pts)
    print(f, 'tabelas dinâmicas', len(pts), names)
    wbx=z.read('xl/workbook.xml').decode('utf8'); print('  abas', len(re.findall('<sheet ',wbx)), 'calcPr', re.search(r'<calcPr[^>]*>',wbx).group(0))
wa=openpyxl.load_workbook(a,read_only=True,data_only=True); wb=openpyxl.load_workbook(b,read_only=True,data_only=True)
assert wa.sheetnames==wb.sheetnames
tot=0
for s in wa.sheetnames:
    if s in RES: continue
    A=wa[s].iter_rows(values_only=True); B=wb[s].iter_rows(values_only=True)
    dif=0; n=0
    import itertools
    for ra,rb in itertools.zip_longest(A,B,fillvalue=()):
        n+=1
        ra=list(ra); rb=list(rb)
        while ra and ra[-1] is None: ra.pop()
        while rb and rb[-1] is None: rb.pop()
        if ra!=rb:
            dif+=1
            if dif<=2: print('   DIF', s, n, [ (j,x,y) for j,(x,y) in enumerate(itertools.zip_longest(ra,rb)) if x!=y][:4])
    tot+=dif
    print(f'{s}: {n} linhas, {dif} diferentes')
print('TOTAL linhas diferentes fora dos resumos:', tot)
