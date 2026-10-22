import openpyxl, pickle, re, json, datetime as dt, collections as C, sys
from openpyxl.utils import get_column_letter as Lt, column_index_from_string as ci
A_, B_ = sys.argv[1], sys.argv[2]
K=pickle.load(open('consolida.pkl','rb')); ex=set(K['excluir']); cel=K['cel']
def load(f, data_only):
    wb=openpyxl.load_workbook(f, read_only=True, data_only=data_only); r=list(wb['SUPERMERCADOS (reuniao)'].iter_rows(values_only=True)); wb.close(); return r
va,vb=load(A_,True),load(B_,True); fa,fb=load(A_,False),load(B_,False)
old=[n for n in range(5,len(va)+1) if any(v is not None for v in va[n-1])]
keep=[n for n in old if n not in ex]
newrows=[n for n in range(5,len(vb)+1) if any(v is not None for v in vb[n-1])]
print('linhas antes', len(old), 'depois', len(newrows), 'esperado', len(keep))
assert len(newrows)==len(keep)
# cabecalho/topo
for i in range(4):
    for j in range(max(len(va[i]),len(vb[i]))):
        x=va[i][j] if j<len(va[i]) else None; y=vb[i][j] if j<len(vb[i]) else None
        if x!=y and not (i==2 and Lt(j+1)=='AJ') and not (i==3 and Lt(j+1)=='AO'): print('TOPO DIF', i+1, Lt(j+1), x, y)
print('AO4:', vb[3][ci('AO')-1], '| AJ3:', va[2][ci('AJ')-1], '->', vb[2][ci('AJ')-1])
norm=lambda f: re.sub(r'(?<=[A-Z])\d+','#',f) if isinstance(f,str) and f.startswith('=') else f
dif=C.Counter(); erros=0; AI={'A':20,'B':15,'C':10,'D':8}; soma=0
for n_old,n_new in zip(keep,newrows):
    ra,rb=va[n_old-1],vb[n_new-1]; ga,gb=fa[n_old-1],fb[n_new-1]
    c=cel.get(n_old,{})
    perm={'L','AO','AJ'}|({'AK','AL','AN'} if 'AK' in c else set())
    for j in range(max(len(ra),len(rb))):
        col=Lt(j+1); x=ra[j] if j<len(ra) else None; y=rb[j] if j<len(rb) else None
        fx=ga[j] if j<len(ga) else None; fy=gb[j] if j<len(gb) else None
        if norm(fx)!=norm(fy) and col not in perm: erros+=1; print('FORMULA', n_old, col, fx, fy) if erros<6 else None
        if x!=y:
            if col in perm: dif[col]+=1
            else: erros+=1; print('DIF', n_old,'->',n_new, col, x, y) if erros<10 else None
    # valores esperados
    Lexp=c.get('L', ra[ci('L')-1])
    if rb[ci('L')-1]!=Lexp: erros+=1; print('L errado', n_old, rb[ci('L')-1], Lexp)
    if rb[ci('AJ')-1]!=AI[rb[ci('J')-1]]*Lexp: erros+=1; print('AJ errado', n_old)
    if 'AK' in c and (rb[ci('AK')-1]!=c['AK'] or rb[ci('AL')-1]!=c['AL'] or rb[ci('AN')-1]!=c['AN']): erros+=1; print('merge errado', n_old)
    if c.get('AO') and rb[ci('AO')-1]!=c['AO']: erros+=1; print('AO errado', n_old)
    soma+=rb[ci('AJ')-1]
print('colunas alteradas:', dict(dif), '| erros:', erros, '| soma AJ', soma)
