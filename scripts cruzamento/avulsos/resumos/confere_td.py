import openpyxl, pickle, collections as C
from openpyxl.utils import column_index_from_string as ci, get_column_letter as L
wb=openpyxl.load_workbook('trabalho.xlsx',read_only=True,data_only=True)
dados={}
def esperado(aba, ini, campos, cst, ccx, fil=None):
    rows=[r for r in list(wb[aba].iter_rows(values_only=True))[ini:] if any(v is not None for v in r)]
    out={}
    for nome,col in campos.items():
        d=C.defaultdict(lambda: C.Counter())
        for r in rows:
            k=r[ci(col)-1]; k=k if k not in (None,'') else '(vazio)'
            s=r[ci(cst)-1]; s={'a tarbalhar':'a trabalhar'}.get(s,s)
            d[k][(s,'n')]+=1; d[k][(s,'cx')]+=r[ci(ccx)-1] or 0
        out[nome]=d
    return out
E={'RESUMO SUPERMERCADO (REUNIAO)':esperado('SUPERMERCADOS (reuniao)',4,{'A':'C','J':'I','S':'AN'},'AL','AJ'),
   '"FARMACIA CONSOLIDADO RESUMO"':esperado('"FARMACIA CONSOLIDADO"',1,{'A':'I','J':'S','S':'R'},'O','T'),
   '"FARMACIA VF RESUMO"':esperado('"FARMACIA VF"',1,{'A':'K','J':'U','S':'V'},'P','W')}
wb.close()
wf=openpyxl.load_workbook('final.xlsx',read_only=True,data_only=True)
for aba,exp in E.items():
    rows=list(wf[aba].iter_rows(values_only=True))
    for col,d in exp.items():
        c0=ci(col)-1
        # acha a linha de cabecalho dos status (linha com 'ativo'/'ATIVO' logo depois do campo)
        h=[i for i,r in enumerate(rows) if len(r)>c0+1 and str(r[c0+1] or '').lower()=='ativo'][0]
        st1, st2 = rows[h][c0+1], rows[h][c0+3]
        assert str(rows[h+1][c0+1])=='Clientes' and str(rows[h+1][c0+2]).startswith('Caixas')
        err=0; n=0; vist=set()
        for r in rows[h+2:]:
            if len(r)<=c0 or r[c0] is None: break
            k=r[c0]
            if str(k)=='Total Geral': tot=r; break
            kk='(vazio)' if str(k).startswith('Sem vendedor') else k
            vist.add(kk)
            ex=d[kk]
            got=[r[c0+1] or 0, r[c0+2] or 0, r[c0+3] or 0, r[c0+4] or 0, r[c0+5], r[c0+6]]
            want=[ex[(st1,'n')], ex[(st1,'cx')], ex[(st2,'n')], ex[(st2,'cx')], ex[(st1,'n')]+ex[(st2,'n')], ex[(st1,'cx')]+ex[(st2,'cx')]]
            n+=1
            if got!=want: err+=1; print('  DIF', aba, col, k, got, want) if err<4 else None
        faltam=set(d)-vist
        print(aba, col, 'status', st1, '|', st2, '- linhas', n, 'divergências', err, 'itens faltando', len(faltam), '| total', tot[c0+1:c0+7])
wf.close()
