"""Prepara a pasta de trabalho: le a aba "I.R.T.T.__Por bandeira (FINAL)" do MAPA MERCADO e as fontes do IRTT, ABRAS e
ranking de farmacias.  Saidas: mapa_irtt.pkl, irtt_layout.xlsx, irtt_grp.pkl, grp_of.pkl, rankings.pkl.
Variaveis: AUTOPEL_MAPA, AUTOPEL_IRTT_DIR (pasta com Ranking-Top300-IRTT.xlsx e IRTT para upload.csv),
AUTOPEL_ABRAS (ABRAS 2026 - CRU.xlsx), AUTOPEL_FARM_RANKING (ranking_lig.pkl de ../farmacias/vf/liga_ranking.py)."""
import os, re, csv, pickle, unicodedata
import openpyxl
MAPA, IRTT_DIR = os.environ["AUTOPEL_MAPA"], os.environ["AUTOPEL_IRTT_DIR"]


def nz(s):
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().upper()
    s = re.sub(r"[\d\*\s]+$", "", re.sub(r"[¹²³⁴⁵]", "", s).strip())
    return re.sub(r"[^A-Z0-9]+", " ", s).strip()


def num(v):  # o IRTT vem com ponto de milhar lido como decimal (1.007 = 1007)
    if isinstance(v, (int, float)):
        return v if float(v).is_integer() else round(v * 1000)
    return None


wb = openpyxl.load_workbook(MAPA, read_only=True)
rows = list(enumerate(wb["I.R.T.T.__Por bandeira (FINAL)"].iter_rows(values_only=True), start=1)); wb.close()
hdr, data = rows[3][1], [(i, r) for i, r in rows[4:] if r[1]]
pickle.dump({"hdr": hdr, "data": data}, open("mapa_irtt.pkl", "wb"))
wo = openpyxl.Workbook(); ws = wo.active  # layout do IRTT (claude) para o irtt.py: sem a coluna F
ws.append([h for k, h in enumerate(hdr[:22]) if k != 5])
for i, r in data: ws.append([v for k, v in enumerate(r[:22]) if k != 5])
wo.save("irtt_layout.xlsx")
wb = openpyxl.load_workbook(os.path.join(IRTT_DIR, "Ranking-Top300-IRTT.xlsx"), read_only=True, data_only=True)
R = list(wb[next(n for n in wb.sheetnames if n.startswith("IRTT (FIM)"))].iter_rows(max_row=400, values_only=True)); wb.close()
grp = {nz(r[4]): dict(pos=r[0], nome=r[4], band=r[5], seg=r[6], venda24=num(r[7]), lojas24=num(r[13]), sede=r[29])
       for r in R[1:] if r[4] and isinstance(r[0], (int, float))}
up = list(csv.DictReader(open(os.path.join(IRTT_DIR, "IRTT para upload.csv"), encoding="utf-8-sig"), delimiter=";"))
pickle.dump({"grp": grp, "up": up}, open("irtt_grp.pkl", "wb"))
pickle.dump({lin: nz(u["rede"]) for (lin, row), u in zip(data, up)}, open("grp_of.pkl", "wb"))  # mesma ordem da aba
wb = openpyxl.load_workbook(os.environ["AUTOPEL_ABRAS"], read_only=True, data_only=True)
ab = [dict(pos=r[0], razao=r[3], uf=r[4], fat=r[5]) for r in list(wb.active.iter_rows(values_only=True))[1:] if r[3] and isinstance(r[5], (int, float))]
wb.close()
pickle.dump({"abras": ab, "farm": pickle.load(open(os.environ["AUTOPEL_FARM_RANKING"], "rb"))}, open("rankings.pkl", "wb"))
print(len(data), "bandeiras |", len(grp), "grupos IRTT |", len(ab), "empresas ABRAS")
