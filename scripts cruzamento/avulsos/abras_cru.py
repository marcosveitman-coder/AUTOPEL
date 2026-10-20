"""Cruza 'ABRAS 2026 - CRU.xlsx' (D = razao social, E = UF) com o faturamento 2019-2026.
Vinculos: CNPJ (do mapa supermercados.xlsx), razao social, nome/marca, produto personalizado e dominio de e-mail
do cadastro (Ultimo faturamento - SQL.csv). Saida: cand_abras.pkl"""
import sys, os, re, glob, pickle, collections
import openpyxl, pandas as pd
S = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
sys.path.insert(0, S)
import common, match
from common import norm_legal, norm

HERE = os.path.dirname(os.path.abspath(__file__))
AB = os.environ.get("AUTOPEL_ABRAS", "ABRAS 2026 - CRU.xlsx")
SMX = os.path.join(common.BASE, "supermercados.xlsx")
CSV = os.path.join(common.BASE, "Ultimo faturamento - SQL.csv")
D = lambda s: re.sub(r"\D", "", str(s or ""))
NS = lambda s: norm_legal(s).replace(" ", "")

# 1) linhas de dados da ABRAS (ignora vazias e TOTAL)
ws = openpyxl.load_workbook(AB, read_only=True).active
data = [(r, row) for r, row in enumerate(ws.iter_rows(values_only=True), start=1)
        if r > 1 and row[3] and str(row[3]).strip().upper() != "TOTAL"]
# CNPJ ja validado no mapa de supermercados (linhas do Ranking ABRAS 2026)
smc = {}
for r in list(openpyxl.load_workbook(SMX, read_only=True).active.iter_rows(values_only=True))[1:]:
    if r[12] == "Ranking ABRAS 2026" and len(D(r[5])) == 14:
        smc.setdefault((NS(r[4]), (r[2] or "").strip()), r[5])
# planilha temporaria no layout que o match.py espera (C=UF, D=fantasia, E=razao, F=CNPJ, H=razao2)
tmp = os.path.join(HERE, "abras_layout.xlsx")
wb = openpyxl.Workbook(); w = wb.active
w.append(["a", "b", "UF", "Nome fantasia", "Razao", "CNPJ", "g", "h"])
for r, row in data:
    uf = (row[4] or "").strip()
    w.append([r, None, uf, row[3], row[3], smc.get((NS(row[3]), uf)), None, None])
wb.save(tmp)


def load_agg_2019():
    root, cli, desc = {}, {}, {}; rows = bad = 0
    for p in sorted(glob.glob(os.path.join(common.AGG, "*.pkl"))):
        if "2018" in os.path.basename(p):
            continue
        d = pickle.load(open(p, "rb"))
        common.merge_into(root, d["root"]); common.merge_into(cli, d["cli"]); common.merge_into(desc, d["desc"])
        rows += d["rows"]; bad += d["bad"]
    return root, cli, desc, rows, bad


match.load_agg = load_agg_2019; match.SM_PATH = tmp; match.CAND = os.path.join(HERE, "cand_abras.pkl")
match.main()
res = pickle.load(open(match.CAND, "rb"))
cand = res["cand"]
root = load_agg_2019()[0]

# 2) dominio de e-mail do cadastro: rotulo do dominio == marca
GENERIC = {"autopel", "gmail", "hotmail", "yahoo", "outlook", "uol", "terra", "bol", "ig", "live", "msn", "icloud",
           "globo", "globomail", "zipmail", "r7", "oi", "aol", "me", "email", "mail"}
df = pd.read_csv(CSV, sep=";", skiprows=1, dtype=str, encoding="utf-8-sig")
lab_roots = collections.defaultdict(set)
lab_uf = collections.defaultdict(set)  # (rotulo, raiz) -> UFs dos CNPJs com esse e-mail
for cnpj, em, uf in zip(df.iloc[:, 2].map(D), df.iloc[:, 8].fillna(""), df.iloc[:, 28].fillna("").str.strip()):
    if len(cnpj) != 14:
        continue
    for e in re.split(r"[;,\s]+", em.lower()):
        m = re.search(r"@([a-z0-9\-]+)\.", e)
        if m and m.group(1) not in GENERIC:
            lab = m.group(1).replace("-", "")
            lab_roots[lab].add(cnpj[:8]); lab_uf[(lab, cnpj[:8])].add(uf)
n_email = 0
for k, (r, row) in enumerate(data):
    keys = set()
    for c in match.brand_cores(row[3]):
        b = c.replace(" ", "").lower()
        if len(b) >= 5:
            keys |= {b, "grupo" + b, "rede" + b, b + "net", "sm" + b, "super" + b, b + "supermercados", b + "supermercado"}
    keys |= {g.lower() for g in match.glued_keys(row[3]) if len(g) >= 5}
    best = None
    for key in keys:
        roots = lab_roots.get(key, set())
        if len(roots) > 15:  # dominio de escritorio/servico, nao de uma rede
            continue
        for rt in roots:
            rec = root.get(rt)
            if not rec or rec[4] not in match.SEG_OK:  # cliente do dominio precisa ser supermercado/alimentacao
                continue
            cc = " " + match.name_core(rec[2]) + " "
            has = any(f" {c} " in cc or c.replace(" ", "") in cc.replace(" ", "") for c in match.brand_cores(row[3]))
            if (row[4] or "").strip() not in lab_uf[(key, rt)]:  # e-mail do dominio precisa estar na UF da rede
                continue
            if best is None or rec[0] > best["dt"]:
                best = match.mk(rec, "email", {"dominio": key, "buyer_has_brand": has})
    if best:
        cand[k]["email"] = best; n_email += 1
print("vinculos por e-mail:", n_email)
pickle.dump({"cand": cand, "rows": [r for r, _ in data], "names": [row[3] for _, row in data],
             "uf": [row[4] for _, row in data], "crm": [row[7] for _, row in data]}, open(match.CAND, "wb"))
