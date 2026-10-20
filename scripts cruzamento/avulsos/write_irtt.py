"""Junta vinculos (CNPJ, razao, codigos da col R, bandeira por nome/produto/e-mail), grava a data na coluna F do IRTT
e confere com a coluna G (dias sem comprar, referencia 13/10/2026)."""
import os, re, sys, glob, pickle, datetime, collections
import openpyxl
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
import common
from choose_irtt import to_date
from irtt import SEG

HERE = os.path.dirname(os.path.abspath(__file__))
IRTT = os.environ.get("AUTOPEL_IRTT", "IRTT (claude).xlsx")
REF = datetime.datetime(2026, 10, 13)
ORDER = ["codigo", "cnpj", "razao", "email", "fantasia", "desc"]
LBL = {"codigo": "Código de cliente (col R)", "cnpj": "CNPJ indicado", "razao": "Razão social indicada",
       "email": "Domínio de e-mail", "fantasia": "Nome da bandeira no cliente", "desc": "Produto personalizado"}
WRITE = "--gravar" in sys.argv

cod = {}
for p in glob.glob(os.path.join(HERE, "cod", "*.pkl")):
    common.merge_into(cod, pickle.load(open(p, "rb")))
R = pickle.load(open(os.path.join(HERE, "cand_irtt.pkl"), "rb"))
import match
from common import norm_legal
AP = lambda s: re.sub(r"['’`´]", "", str(s or ""))


def sig(text):
    """palavras significativas (4+ letras, nao genericas) + forma sem espaco"""
    toks = [t for t in norm_legal(AP(text)).split() if len(t) >= 4 and t not in match.WEAK]
    return set(toks), "".join(norm_legal(AP(text)).split())


def related(rec, row):
    if len(rec[3]) == 14 and re.sub(r"\D", "", str(row[19] or ""))[:8] == rec[3][:8]:
        return True
    ct, cns = sig(rec[2])
    for ref in (row[0], row[1], row[20]):
        rt, rns = sig(ref)
        if ct & rt or any(len(t) >= 5 and t in cns for t in rt):
            return True
    return False
for k, row in enumerate(R["data"]):
    best = None
    for c in re.findall(r"\d+", str(row[17] or "")):
        rec = cod.get(c.zfill(6))
        if rec and row[2] != "Outros Segmentos" and not related(rec, row) and rec[4] and rec[4] not in SEG.get(row[2], set()):
            continue  # codigo usado por empresa sem relacao e de outro segmento (ex.: codigo reaproveitado por outra empresa)
        if rec and (best is None or rec[0] > best["dt"]):
            best = {"dt": rec[0], "vends": dict(rec[1]), "cliente": rec[2], "cnpj": rec[3], "seg": rec[4], "how": "codigo"}
    if best:
        R["cand"][k]["codigo"] = best


def choose(c):
    strong = [(m, c[m]) for m in ("codigo", "cnpj", "razao") if m in c]
    weak = [(m, c[m]) for m in ("email", "fantasia", "desc") if m in c]
    pick = lambda l: max(l, key=lambda md: (md[1]["dt"], -ORDER.index(md[0])))
    if not strong:
        return pick(weak) if weak else None
    best = pick(strong)
    over = [md for md in weak if md[1].get("buyer_has_brand") and md[1]["dt"] > best[1]["dt"]]
    return pick(over) if over else best


st = collections.Counter(); diverg = []; yrs = collections.Counter(); res = {}
for k, c in enumerate(R["cand"]):
    r, row = R["rows"][k], R["data"][k]
    ch = choose(c)
    res[r] = ch
    if ch:
        st[LBL[ch[0]]] += 1
        yrs[to_date(ch[1]["dt"]).year] += 1
    g = row[6]
    if g is None:
        st["G vazio | " + ("nós achamos" if ch else "ninguém tem")] += 1
        continue
    if not ch:
        st["G existe, nós não achamos"] += 1; diverg.append((r, row[1][:30], "-", g, "")); continue
    ours = (REF - to_date(ch[1]["dt"])).days
    tag = "bate com G (±3 dias)" if abs(ours - g) <= 3 else ("nossa MAIS RECENTE" if ours < g else "nossa MAIS ANTIGA")
    st[tag] += 1
    if tag != "bate com G (±3 dias)":
        diverg.append((r, row[1][:30], to_date(ch[1]["dt"]).date(), f"G={g} nosso={ours}", f'{ch[0]}: {ch[1]["cliente"][:32]}'))
print({k: v for k, v in sorted(st.items())})
print("por ano:", dict(sorted(yrs.items())))
print("\nDIVERGENCIAS:")
for d in diverg:
    print("  ", d)

if WRITE:
    try:
        open(IRTT, "r+b").close()
    except PermissionError:
        raise SystemExit("BLOQUEADO - feche o IRTT no Excel")
    wb = openpyxl.load_workbook(IRTT); ws = wb.active
    for r, ch in res.items():
        if ch:
            x = ws.cell(r, 6, to_date(ch[1]["dt"])); x.number_format = "DD/MM/YYYY"
        else:
            ws.cell(r, 6, None)
    wb.save(IRTT)
    print("gravado:", sum(1 for ch in res.values() if ch), "datas na coluna F")
