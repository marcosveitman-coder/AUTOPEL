"""Monta a planilha de rotulos vendidos a partir do rotulos.pkl (busca_rotulos.py).
Uso: python planilha_rotulos.py rotulos.pkl saida.xlsx"""
import sys, pickle, collections, datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
R = pickle.load(open(sys.argv[1], "rb"))
v = lambda x: float(x) if isinstance(x, (int, float)) else 0.0


def grupo(r):
    d = str(r["DES_PROD"]).upper()
    if r["TIPO"].startswith("ROTULO"):
        return "Fita para rotulador" if "FITA" in d or "ROTULADOR" in d else "Rótulo (palavra na descrição)"
    return "Etiqueta BOPP (rótulo adesivo)"


for r in R: r["GRUPO"] = grupo(r); r["ANO"] = (r["EMISSAO"] or 0) // 10000
GR = ["Rótulo (palavra na descrição)", "Etiqueta BOPP (rótulo adesivo)", "Fita para rotulador"]
ANOS = sorted({r["ANO"] for r in R})
wb = openpyxl.Workbook()
H = Font(bold=True, color="FFFFFF"); HF = PatternFill("solid", fgColor="1F4E78")


def cab(ws, cols, larg):
    ws.append(cols)
    for j, w in enumerate(larg, 1):
        c = ws.cell(1, j); c.font = H; c.fill = HF; c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A2"


ws = wb.active; ws.title = "Resumo"
cab(ws, ["Grupo", *[str(a) for a in ANOS], "Total R$", "Linhas", "Clientes (CNPJ)", "Produtos"], [34] + [12] * len(ANOS) + [14, 9, 10, 10])
for g in GR:
    rs = [r for r in R if r["GRUPO"] == g]
    ws.append([g, *[round(sum(v(r["VLR_TOTAL"]) for r in rs if r["ANO"] == a), 2) for a in ANOS],
               round(sum(v(r["VLR_TOTAL"]) for r in rs), 2), len(rs), len({r["CNPJ"] for r in rs}), len({r["DES_PROD"] for r in rs})])
for row in ws.iter_rows(min_row=2, min_col=2, max_col=len(ANOS) + 2):
    for c in row: c.number_format = "#,##0"

ws = wb.create_sheet("Por produto")
cab(ws, ["Grupo", "Cód. produto", "Descrição", "Qtd", "Valor R$", "Clientes", "1ª venda", "Última venda"], [30, 14, 60, 10, 14, 9, 12, 12])
pp = collections.defaultdict(list)
for r in R: pp[(r["GRUPO"], str(r["COD_PROD"] or "").strip(), str(r["DES_PROD"]).strip())].append(r)
toD = lambda d: datetime.date(d // 10000, d // 100 % 100, d % 100) if d else None
for (g, cp, dp), rs in sorted(pp.items(), key=lambda kv: (GR.index(kv[0][0]), -sum(v(r["VLR_TOTAL"]) for r in kv[1]))):
    ds = [r["EMISSAO"] for r in rs if r["EMISSAO"]]
    ws.append([g, cp, dp, sum(v(r["QUANT"]) for r in rs), round(sum(v(r["VLR_TOTAL"]) for r in rs), 2), len({r["CNPJ"] for r in rs}),
               toD(min(ds)) if ds else None, toD(max(ds)) if ds else None])

ws = wb.create_sheet("Por cliente")
cab(ws, ["Grupo", "Cliente", "CNPJ", "UF", "Segmento", "Valor R$", "Linhas", "Última compra", "Vendedor da última venda"], [30, 42, 20, 5, 18, 14, 8, 12, 34])
pc = collections.defaultdict(list)
for r in R: pc[(r["GRUPO"], str(r["CNPJ"] or "").strip())].append(r)
for (g, cnpj), rs in sorted(pc.items(), key=lambda kv: (GR.index(kv[0][0]), -sum(v(r["VLR_TOTAL"]) for r in kv[1]))):
    u = max(rs, key=lambda r: (r["EMISSAO"] or 0, v(r["VLR_TOTAL"])))
    ws.append([g, str(u["CLIENTE"] or "").strip(), cnpj, str(u["ESTADO"] or "").strip(), str(u["SEGMENTO"] or "").strip(),
               round(sum(v(r["VLR_TOTAL"]) for r in rs), 2), len(rs), toD(u["EMISSAO"]), str(u["NOME_VENDEDOR"] or "").strip()])

ws = wb.create_sheet("Linhas")
cols = ["GRUPO", "EMISSAO", "COD_CLI", "COD_LOJA", "CLIENTE", "CNPJ", "ESTADO", "SEGMENTO", "NOME_VENDEDOR", "COD_PROD", "DES_PROD", "QUANT", "VLR_TOTAL", "DESC_FAMILIA", "ARQUIVO"]
cab(ws, cols, [28, 11, 9, 8, 40, 20, 5, 18, 32, 14, 60, 9, 12, 26, 30])
for r in sorted(R, key=lambda r: (GR.index(r["GRUPO"]), r["EMISSAO"] or 0)):
    ws.append([toD(r[c]) if c == "EMISSAO" else (str(r[c]).strip() if isinstance(r[c], str) else r[c]) for c in cols])
for name, cs in (("Por produto", "GH"), ("Por cliente", "H"), ("Linhas", "B")):
    for c in cs:
        for cell in wb[name][c][1:]: cell.number_format = "DD/MM/YYYY"
for name, cs in (("Por produto", "E"), ("Por cliente", "F"), ("Linhas", "M")):
    for c in cs:
        for cell in wb[name][c][1:]: cell.number_format = "#,##0.00"
for w in wb.worksheets: w.auto_filter.ref = w.dimensions
wb.save(sys.argv[2]); print("ok", len(R))
