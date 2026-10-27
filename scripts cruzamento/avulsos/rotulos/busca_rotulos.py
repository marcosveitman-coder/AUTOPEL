"""Procura rotulos nas descricoes de produto (DES_PROD) de todas as bases fat AAAA.N.xlsx.
Pega: ROTUL* (rotulo, rotulos, rotul., rotulador), ETIQUETA BOPP (material de rotulo adesivo).
A aba Planilha1 de fat 2025.2 e copia de dez/2025 da Listagem e fica fora.  Saida: rotulos.pkl (lista de linhas)."""
import os, sys, re, glob, pickle, time, unicodedata
from multiprocessing import Pool
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", ".."))
from common import BASE
from extract import to_int_date
N = lambda s: unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().upper()
RX_ROT = re.compile(r"\bROTUL")
RX_BOPP = re.compile(r"\bETIQ\w*\b.*\bBOPP\b|\bBOPP\b.*\bETIQ")
COLS = ["EMISSAO", "COD_CLI", "COD_LOJA", "NOME_VENDEDOR", "CNPJ", "ESTADO", "CLIENTE", "COD_PROD", "DES_PROD", "QUANT",
        "VLR_TOTAL", "SEGMENTO", "DESC_FAMILIA"]


def work(job):
    f, sh = job
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(BASE, f), read_only=True)
    it = wb[sh].iter_rows(values_only=True)
    hdr = [str(h or "").strip().upper() for h in next(it, ())]
    idx = {c: hdr.index(c) if c in hdr else None for c in COLS}
    out = []
    for r in it:
        d = N(r[idx["DES_PROD"]]) if idx["DES_PROD"] is not None else ""
        tipo = "ROTULO (palavra na descrição)" if RX_ROT.search(d) else "ETIQUETA BOPP" if RX_BOPP.search(d) else None
        if not tipo:
            continue
        x = {c: (r[i] if i is not None and i < len(r) else None) for c, i in idx.items()}
        x["EMISSAO"] = to_int_date(x["EMISSAO"]); x["TIPO"] = tipo; x["ARQUIVO"] = f + " / " + sh
        out.append(x)
    wb.close()
    return f + " " + sh, out


if __name__ == "__main__":
    import openpyxl
    jobs = []
    for f in sorted(os.path.basename(p) for p in glob.glob(os.path.join(BASE, "fat 20*.xlsx"))):
        wb = openpyxl.load_workbook(os.path.join(BASE, f), read_only=True)
        jobs += [(f, s) for s in wb.sheetnames if not (f.startswith("fat 2025.2") and s == "Planilha1")]
        wb.close()
    t = time.time(); rows = []
    with Pool(4) as p:
        for name, out in p.imap_unordered(work, jobs):
            rows += out; print(name, len(out), f"{time.time()-t:.0f}s", flush=True)
    pickle.dump(rows, open(sys.argv[1] if len(sys.argv) > 1 else "rotulos.pkl", "wb"))
    print("TOTAL linhas:", len(rows))
