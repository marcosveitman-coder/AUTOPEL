"""Tira das bases de faturamento todas as linhas de uma ou mais raizes de CNPJ (todas as colunas).

Uso: python -I linhas_cliente.py SAIDA.json RAIZ[,RAIZ...] "fat 2020.1.xlsx" ["fat 2022.1.xlsx" ...]
Os arquivos ficam na pasta "Bases Claude"; le todas as abas, menos a Planilha1 de fat 2025.2 (copia de dez/2025).
Colunas pelo nome do cabecalho.
"""
import json
import os
import re
import sys
from multiprocessing import Pool

# pasta "Bases Claude" = tres niveis acima (avulsos/kraft_heinz/ -> scripts cruzamento/ -> Bases Claude)
BASE = os.environ.get("AUTOPEL_BASES", os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", ".."))
NAO_DIGITO = re.compile(r"\D")


def trabalho(tarefa):
    arquivo, raizes = tarefa
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(BASE, arquivo), read_only=True)
    achadas = []
    for aba in wb.sheetnames:
        if arquivo.startswith("fat 2025.2") and aba == "Planilha1":
            continue
        it = wb[aba].iter_rows(values_only=True)
        cab = [str(h).strip() if h is not None else "" for h in next(it, ())]
        if "CNPJ" not in cab:
            continue
        i_cnpj = cab.index("CNPJ")
        for linha in it:
            if i_cnpj < len(linha) and NAO_DIGITO.sub("", str(linha[i_cnpj] or ""))[:8] in raizes:
                achadas.append({"arquivo": arquivo, "aba": aba,
                                **{c: (str(v) if v is not None else None) for c, v in zip(cab, linha) if c}})
    wb.close()
    return achadas


if __name__ == "__main__":
    saida, raizes, arquivos = sys.argv[1], set(sys.argv[2].split(",")), sys.argv[3:]
    with Pool(min(4, len(arquivos))) as pool:
        linhas = [l for parte in pool.map(trabalho, [(a, raizes) for a in arquivos]) for l in parte]
    json.dump(linhas, open(saida, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print(len(linhas), "linhas")
