"""Ultima venda por COD_CLI (bases fat 2019-2026). Saida: cod/<arquivo>__<aba>.pkl {cod: [data, {vend: valor}, cliente, cnpj, seg]}"""
import os, sys, glob, pickle, time, re
from multiprocessing import Pool
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from common import BASE
from extract import to_int_date, upd

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "cod")
FILES = sorted(os.path.basename(p) for p in glob.glob(os.path.join(BASE, "fat *.xlsx")) if "2018" not in p)
NEED = ["EMISSAO", "COD_CLI", "NOME_VENDEDOR", "CNPJ", "CLIENTE", "VLR_TOTAL", "SEGMENTO"]


def work(job):
    fname, sheet = job
    import openpyxl
    t = time.time()
    wb = openpyxl.load_workbook(os.path.join(BASE, fname), read_only=True)
    it = wb[sheet].iter_rows(values_only=True)
    hdr = [str(h).strip() if h is not None else "" for h in next(it, ())]
    if any(n not in hdr for n in NEED):
        wb.close()
        return f"{fname} [{sheet}] IGNORADA"
    iE, iK, iV, iC, iCl, iT, iS = (hdr.index(n) for n in NEED)
    cod, n = {}, 0
    for r in it:
        dt = to_int_date(r[iE])
        k = str(r[iK] or "").strip()
        if dt is None or not k:
            continue
        n += 1
        try:
            val = float(r[iT] or 0)
        except (TypeError, ValueError):
            val = 0.0
        vend = str(r[iV] or "").strip()
        upd(cod, k.zfill(6), dt, vend, val, str(r[iCl] or "").strip(), re.sub(r"\D", "", str(r[iC] or "")),
            str(r[iS] or "").strip())
    wb.close()
    pickle.dump(cod, open(os.path.join(OUT, f"{fname}__{sheet}.pkl"), "wb"))
    return f"{fname} [{sheet}] linhas={n} codigos={len(cod)} t={time.time()-t:.0f}s"


if __name__ == "__main__":
    import openpyxl
    os.makedirs(OUT, exist_ok=True)
    js = []
    for f in FILES:
        wb = openpyxl.load_workbook(os.path.join(BASE, f), read_only=True)
        js += [(f, s) for s in wb.sheetnames]
        wb.close()
    js.sort(key=lambda j: -os.path.getsize(os.path.join(BASE, j[0])))
    print("abas:", js, flush=True)
    t = time.time()
    with Pool(4) as p:
        for msg in p.imap_unordered(work, js):
            print(msg, flush=True)
    print(f"TOTAL {time.time()-t:.0f}s", flush=True)
