"""Saude do estoque (IQR, slow, phase out), triagem do excedente para e-commerce e concentracao de clientes,
calculadas com os dados da planilha de compras (aba Geral, colunas do TOTVS + colunas CLAUDE ja calculadas).

    python saude_ecommerce.py PLANILHA.xlsx PASTA_PKL "Ultimo faturamento - SQL.csv" SAIDA_DIR [AAAA-MM-DD]

PASTA_PKL tem os arquivos dos outros scripts: demanda.pkl e dup_planilha1.pkl (extrai_demanda.py: venda por mes e por
cliente nas bases de faturamento) e demanda_totvs.pkl (venda mensal do TOTVS, todos os canais, das fotos diarias).
O CSV do ultimo faturamento so traduz COD_CLI -> raiz do CNPJ e nome. Saida:
- saude.json      : tabelas das secoes 7, 8 e 9 da aba "analise compras CLAUDE" (grava_saude.ps1);
- param_saude.csv : colunas novas por SKU da tabela ParamClaude.

Regras:
- Venda recente = venda no TOTVS (todos os canais, abr-set) ou nas bases de faturamento nos ultimos 6 meses.
  Sem venda em 6 meses: slow se as bases mostram venda nos 12 meses; senao, phase out potencial.
- Estoque comprometido (Reserva + VendProt, ate o fisico) conta como saudavel.
- Saudavel estrito = comprometido + alvo CLAUDE do mes vigente; ampliado = + 2 meses da previsao CLAUDE.
  O que passa do ampliado e excesso. Valor rateado pelo custo medio do estoque (Vlr Estoque / Fisico).
- Sem fisico, sem valor ou valor negativo = nao classificado (valor negativo fica fora do total).
- Excedente livre = excesso + slow + phase out (nao comprometido). Triagem: personalizado (nome de cliente na
  descricao, PERSONALIZ/LOGO, grupo 0015), concentrado (1 cliente com 80%+ em 12 meses), revisar requisitos
  (alimento, EPI, saude, eletrico/volumoso) e generico (candidato).
- Concentracao: quantidade por cliente (raiz do CNPJ) de out/25 a set/26 nas bases, sem a copia de dez/25.
"""
import collections
import csv
import datetime as dt
import json
import pickle
import re
import sys
from pathlib import Path

import openpyxl

GRUPOS = {'0001': 'Informática e eletrônicos', '0002': 'Papelaria e escritório', '0003': 'Bobinas',
          '0004': 'Limpeza e higiene', '0005': 'Papéis', '0006': 'Etiquetas e ribbons',
          '0013': 'Mobiliário e eletro', '0015': 'Personalizados', '0016': 'EPI', '0017': 'Copa e cozinha'}
REQUISITO = re.compile(r'\b(CAFE|ACUCAR|ADOCANTE|CHA|BISCOITO|LEITE|SUCO|AGUA MINERAL|ACHOCOLATADO|GLICEMIA|MEDIC|'
                       r'LUVA|MASCARA|ALCOOL|TERMOMETRO|PRESSAO|SERINGA|AGULHA|REFRIGERADOR|FRIGOBAR|FREEZER|'
                       r'BEBEDOURO|MICROONDAS|FOGAO|CAFETEIRA|VENTILADOR|PURIFICADOR|CADEIRA|ARMARIO|MESA)\b')
PERSONAL = re.compile(r'PERSONALIZ|\bLOGO\b|\bLOGOMARCA\b|\bEXCLUSIV')
GENERICO_NOME = set('''COMERCIO COMERCIAL LTDA EIRELI SERVICOS SERVICO INDUSTRIA INDUSTRIAL DISTRIBUIDORA DISTRIBUIDOR
BRASIL BRASILEIRA SUPERMERCADO SUPERMERCADOS ATACADO ATACADISTA PAPELARIA EMPREENDIMENTOS PARTICIPACOES GRUPO LOJAS
EMPRESA NACIONAL PAULO PAULISTA FARMACIA FARMACIAS DROGARIA DROGARIAS HOSPITAL ESTADO MUNICIPIO PREFEITURA SECRETARIA
GOVERNO FEDERAL BANCO CAIXA ECONOMICA ASSOCIACAO INSTITUTO FUNDACAO CENTRO SOCIEDADE ADMINISTRACAO TECNOLOGIA SISTEMAS
SOLUCOES LOGISTICA TRANSPORTES TRANSPORTE IMPORTACAO EXPORTACAO PRODUTOS ALIMENTOS MATERIAIS MATERIAL EQUIPAMENTOS
ESCRITORIO LIMPEZA HIGIENE PLASTICOS EMBALAGENS PAPEIS DIVERSOS CLIENTES GERAL GERAIS UNIDADE MATRIZ FILIAL SHOPPING
CENTER PAPEL CAIXA SACOLA BRANCO BRANCA PRETO PRETA AZUL VERDE AMARELO VERMELHO TRANSPARENTE PLASTICA PLASTICO'''.split())


def num(x):
    try:
        return float(x)
    except (TypeError, ValueError):
        return 0.0


def meses_ate(hoje, n):
    """Lista dos n meses que terminam no mes de hoje (inclusive), AAAA-MM."""
    a, m = hoje.year, hoje.month
    out = []
    for _ in range(n):
        out.append(f'{a:04d}-{m:02d}')
        m -= 1
        if m == 0:
            a, m = a - 1, 12
    return out


def le_geral(arq):
    wb = openpyxl.load_workbook(arq, read_only=True, data_only=True)
    ws = wb[next(n for n in wb.sheetnames if n.startswith('Geral (arq.') and n.endswith(')'))]
    rows = ws.iter_rows(values_only=True)
    next(rows)
    H = [str(x).strip() if x is not None else '' for x in next(rows)]
    ix = {h: i for i, h in enumerate(H) if h}
    dados = [r for r in rows if r and r[ix['Codigo']]]
    wb.close()
    return ix, dados


def clientes(csv_path):
    """COD_CLIENT -> (raiz do CNPJ ou codigo, nome)."""
    out = {}
    with open(csv_path, encoding='utf-8-sig', errors='replace') as f:
        rd = csv.reader(f, delimiter=';')
        next(rd), next(rd)
        for r in rd:
            if len(r) < 7:
                continue
            cod = r[0].strip()
            if cod in out:
                continue
            doc = re.sub(r'\D', '', r[2])
            out[cod] = (doc[:8] if len(doc) == 14 else "C" + cod, (r[5].strip() or r[6].strip()), r[6].strip())
    return out


def tokens_nome(*nomes):
    t = set()
    for n in nomes:
        for w in re.findall(r'[A-Z0-9]+', (n or '').upper()):
            if len(w) >= 5 and not w.isdigit() and w not in GENERICO_NOME:
                t.add(w)
    return t


def main(arq, pasta, csv_cli, saida, hoje):
    P, S = Path(pasta), Path(saida)
    S.mkdir(parents=True, exist_ok=True)
    ix, dados = le_geral(arq)
    dem = pickle.load(open(P / 'demanda.pkl', 'rb'))
    dup = pickle.load(open(P / 'dup_planilha1.pkl', 'rb'))
    tv = pickle.load(open(P / 'demanda_totvs.pkl', 'rb'))['tot']
    cad = clientes(csv_cli)

    m6, m12 = set(meses_ate(hoje, 6)), set(meses_ate(hoje, 12))
    ult_fat = {}
    fat_6m = collections.defaultdict(float)
    for (fil, cod, mes), (q, v, n) in dem['mes'].items():
        if q > 0 and mes > ult_fat.get(cod, ''):
            ult_fat[cod] = mes
        if mes in m6 and mes < hoje.strftime('%Y-%m'):
            fat_6m[cod] += q
    cli = collections.defaultdict(lambda: collections.defaultdict(float))
    for (fil, cod, cc), q in dem['cli12'].items():
        q -= dup['cli12'].get((fil, cod, cc), 0.0)
        if q > 0:
            cli[cod][cad.get(cc, ('C' + cc, cc, cc))[0]] += q
    nome_raiz = {}
    for cc, (raiz, fant, nome) in cad.items():
        nome_raiz.setdefault(raiz, (fant, nome))

    linhas, sku = [], []
    for r in dados:
        g = lambda n: r[ix[n]]
        cod = str(g('Codigo')).strip()
        fis, vlr = num(g('Fisico')), num(g('Vlr Estoque'))
        res, vp = max(0.0, num(g('Reserva'))), max(0.0, num(g('VendProt')))
        alvo, prev = max(0.0, num(g('CLAUDE alvo mês vigente'))), max(0.0, num(g('CLAUDE previsão mês vigente QTD')))
        tvm = tv.get(cod, {})
        tv_meses = [m for m, q in tvm.items() if q > 0]
        ult = max(tv_meses + ([ult_fat[cod]] if cod in ult_fat else []), default='')
        recente = bool(ult) and ult >= min(m6)
        doze = bool(ult) and ult >= min(m12)
        # concentracao
        cq = cli.get(cod, {})
        tot = sum(cq.values())
        ncli = len(cq)
        top_raiz, top_q = max(cq.items(), key=lambda kv: kv[1]) if cq else ('', 0.0)
        top_share = top_q / tot if tot else None
        hhi = sum((q / tot) ** 2 for q in cq.values()) if tot else None
        tv6 = sum(q for m, q in tvm.items() if m in m6)
        cobert = fat_6m.get(cod, 0.0) / tv6 if tv6 > 0 else None
        classe_cli = ('' if not ncli else '1 cliente' if ncli == 1 else '2 clientes' if ncli == 2
                      else '3 a 4 clientes' if ncli <= 4 else '5 ou mais')
        top_nome = nome_raiz.get(top_raiz, ('', ''))[0] if top_raiz else ''
        # saude
        com = min(max(fis, 0.0), res + vp)
        unit = vlr / fis if fis > 0 and vlr > 0 else 0.0
        q = dict(estrito=0.0, amplo=0.0, excesso=0.0, slow=0.0, phase=0.0)
        if fis <= 0 or vlr <= 0:
            classe = 'Não classificado'
        elif recente:
            lim_e, lim_a = com + alvo, com + alvo + 2 * prev
            q['estrito'] = min(fis, lim_e)
            q['amplo'] = min(fis, lim_a)
            q['excesso'] = fis - q['amplo']
            classe = 'Excesso' if q['excesso'] > q['amplo'] else 'Saudável'
        else:
            q['estrito'] = q['amplo'] = com
            livre = fis - com
            if doze:
                q['slow'], classe = livre, 'Slow moving'
            else:
                q['phase'], classe = livre, 'Phase out potencial'
        rs = {k: v * unit for k, v in q.items()}
        nao_class = vlr if classe == 'Não classificado' and vlr > 0 else 0.0
        livre_q = q['excesso'] + q['slow'] + q['phase']
        livre_rs = livre_q * unit
        # triagem e-commerce
        desc = str(g('Descricao') or '').upper()
        grupo = str(g('Grupo') or '').strip()
        nomes_cli = set()
        for raiz in sorted(cq, key=cq.get, reverse=True)[:5]:
            nomes_cli |= tokens_nome(*nome_raiz.get(raiz, ('', '')))
        palavras = set(re.findall(r'[A-Z0-9]+', desc))
        if livre_q <= 0:
            triagem = ''
        elif PERSONAL.search(desc) or grupo == '0015' or (nomes_cli & palavras):
            triagem = 'Personalizado / cliente específico'
        elif top_share is not None and top_share >= 0.8 and tot > 0:
            triagem = 'Concentrado em 1 cliente'
        elif grupo in ('0016', '0017') or REQUISITO.search(desc):
            triagem = 'Revisar requisitos (validade, CA, elétrico, volume)'
        else:
            triagem = 'Genérico: candidato a e-commerce'
        precos = [num(g(c)) for c in ('Oceano_Preco', 'Kalunga_Preco', 'Mercado Livre_Preco', 'Kalunga_Preco (FOGELMAN)')]
        precos = [p for p in precos if p > 0]
        custo = num(g('Ult.Preco')) or unit
        # preco de mercado so vale perto do custo: fora de 0,3x a 5x o produto comparado e outro (ou outra embalagem)
        precos = [p for p in precos if custo and 0.3 <= p / custo <= 5]
        pm = min(precos) if precos else None
        sku.append(dict(cod=cod, comprador=str(g('Comprador') or '').strip(), desc=str(g('Descricao') or '').strip(),
                        familia=str(g('FAMILIA') or '').strip(), grupo=GRUPOS.get(grupo, 'Outros'), abc=str(g('CLAUDE classe (ABC-XYZ-PQR)') or '')[:1],
                        vlr=vlr, fis=fis, ult=ult, classe=classe, rs=rs, nao_class=nao_class, livre_q=livre_q,
                        livre_rs=livre_rs, triagem=triagem, preco_merc=pm, custo=custo, ncli=ncli,
                        classe_cli=classe_cli, top_share=top_share, top_nome=top_nome,
                        cli_equiv=(1 / hhi if hhi else None), cobert=cobert,
                        compra3=sum(num(g(c)) for c in ('CLAUDE mês vigente R$', 'CLAUDE próximo mês R$', 'CLAUDE mês subsequente R$'))))

    # --- tabelas -------------------------------------------------------------------------------------------------
    tot_pos = sum(s['vlr'] for s in sku if s['vlr'] > 0)
    soma = lambda f, lst=sku: sum(f(s) for s in lst)
    estrito = soma(lambda s: s['rs']['estrito'])
    amplo = soma(lambda s: s['rs']['amplo'])
    exc, slow, phase = soma(lambda s: s['rs']['excesso']), soma(lambda s: s['rs']['slow']), soma(lambda s: s['rs']['phase'])
    nc = soma(lambda s: s['nao_class'])
    cnt = collections.Counter(s['classe'] for s in sku)
    resumo = [['Parcela do estoque (custo)', 'SKUs (pela parcela maior)', 'R$', '% do estoque'],
              ['Saudável: comprometido + alvo CLAUDE (IQR estrito)', None, estrito, estrito / tot_pos],
              ['Saudável ampliado: + 2 meses de previsão (IQR ampliado)', cnt['Saudável'], amplo, amplo / tot_pos],
              ['Excesso: acima do ampliado, com venda recente', cnt['Excesso'], exc, exc / tot_pos],
              ['Slow moving: sem venda há 6 meses, vendeu nos 12', cnt['Slow moving'], slow, slow / tot_pos],
              ['Phase out potencial: sem venda conhecida há 12 meses', cnt['Phase out potencial'], phase, phase / tot_pos],
              ['Não classificado: sem físico ou sem valor', cnt['Não classificado'], nc, nc / tot_pos],
              ['Total', len(sku), amplo + exc + slow + phase + nc, (amplo + exc + slow + phase + nc) / tot_pos]]

    def por(chave, titulo, ordem=None):
        gr = collections.defaultdict(list)
        for s in sku:
            gr[s[chave]].append(s)
        ks = ordem or sorted(gr, key=lambda k: -soma(lambda s: max(0, s['vlr']), gr[k]))
        t = [[titulo, 'SKUs', 'Estoque R$', 'IQR estrito', 'IQR ampliado', 'Excesso R$', 'Slow R$', 'Phase out R$', 'Não classificado R$']]
        for k in ks:
            l = gr.get(k, [])
            if not l:
                continue
            v = soma(lambda s: max(0, s['vlr']), l)
            t.append([k or '(vazio)', len(l), v, soma(lambda s: s['rs']['estrito'], l) / v if v else None,
                      soma(lambda s: s['rs']['amplo'], l) / v if v else None, soma(lambda s: s['rs']['excesso'], l),
                      soma(lambda s: s['rs']['slow'], l), soma(lambda s: s['rs']['phase'], l), soma(lambda s: s['nao_class'], l)])
        t.append(['Total', len(sku), tot_pos, estrito / tot_pos, amplo / tot_pos, exc, slow, phase, nc])
        return t

    top_exc = sorted(sku, key=lambda s: -s['livre_rs'])[:30]
    t_top = [['Código', 'Descrição', 'Comprador', 'Saúde', 'Última venda', 'Estoque R$', 'Fora da política R$', 'Triagem']]
    for s in top_exc:
        t_top.append([s['cod'], s['desc'][:60], s['comprador'], s['classe'], s['ult'] or 'sem venda conhecida', s['vlr'], s['livre_rs'], s['triagem']])

    ordem_tri = ['Genérico: candidato a e-commerce', 'Revisar requisitos (validade, CA, elétrico, volume)',
                 'Concentrado em 1 cliente', 'Personalizado / cliente específico']
    t_tri = [['Triagem do excedente livre', 'SKUs', 'Excedente livre R$', '% do excedente', 'Com preço de mercado na Geral', 'O que fazer']]
    acao = {ordem_tri[0]: 'Lote piloto no e-commerce/marketplace, começando pelos de maior valor e com preço de mercado acima do custo.',
            ordem_tri[1]: 'Conferir validade, certificado (CA) e frete; vender por canal próprio ou devolver ao fornecedor.',
            ordem_tri[2]: 'Negociar com o cliente (antecipar pedido, contrato) antes de oferecer a terceiros.',
            ordem_tri[3]: 'Não vender a terceiros: cobrar consumo do cliente dono, combinar contrato ou destruição.'}
    tot_livre = soma(lambda s: s['livre_rs'])
    for k in ordem_tri:
        l = [s for s in sku if s['triagem'] == k]
        t_tri.append([k, len(l), soma(lambda s: s['livre_rs'], l), soma(lambda s: s['livre_rs'], l) / tot_livre,
                      sum(1 for s in l if s['preco_merc']), acao[k]])
    t_tri.append(['Total', sum(r[1] for r in t_tri[1:]), tot_livre, 1.0, sum(r[4] for r in t_tri[1:]), ''])

    gen = sorted([s for s in sku if s['triagem'] == ordem_tri[0]], key=lambda s: -s['livre_rs'])[:40]
    t_gen = [['Código', 'Descrição', 'Família', 'Comprador', 'Excedente QTD', 'Custo unitário R$', 'Excedente R$',
              'Menor preço de mercado (Geral)', 'Preço de mercado ÷ custo', 'Última venda']]
    for s in gen:
        t_gen.append([s['cod'], s['desc'][:60], s['familia'], s['comprador'], round(s['livre_q']), s['custo'], s['livre_rs'],
                      s['preco_merc'], (s['preco_merc'] / s['custo'] if s['preco_merc'] and s['custo'] else None), s['ult'] or 'sem venda conhecida'])

    ordem_cli = ['1 cliente', '2 clientes', '3 a 4 clientes', '5 ou mais', '']
    t_cli = [['Clientes em 12 meses (bases de faturamento)', 'SKUs', 'Estoque R$', 'Compra CLAUDE out–dez R$', 'Excedente livre R$',
              'SKUs com 1 cliente ≥ 80%', 'Regime sugerido']]
    reg = {'1 cliente': 'Sob encomenda ou contrato com o cliente (não manter estoque sem compromisso).',
           '2 clientes': 'Previsão com o comercial; estoque só com pedido ou contrato firme.',
           '3 a 4 clientes': 'Reposição com lote menor e revisão com o comercial.',
           '5 ou mais': 'Reposição normal pela metodologia CLAUDE.',
           '': 'Sem venda nas bases em 12 meses (ou canal fora das bases): usar a venda do TOTVS.'}
    for k in ordem_cli:
        l = [s for s in sku if s['classe_cli'] == k]
        t_cli.append([k or 'Sem cliente nas bases', len(l), soma(lambda s: max(0, s['vlr']), l), soma(lambda s: s['compra3'], l),
                      soma(lambda s: s['livre_rs'], l), sum(1 for s in l if s['top_share'] is not None and s['top_share'] >= 0.8), reg[k]])
    t_cli.append(['Total', len(sku), tot_pos, soma(lambda s: s['compra3']), tot_livre, sum(r[5] for r in t_cli[1:]), ''])

    conc = sorted([s for s in sku if s['top_share'] is not None and s['top_share'] >= 0.8 and (s['compra3'] > 0 or s['vlr'] > 0)],
                  key=lambda s: -(s['compra3'] + max(0, s['vlr'])))[:30]
    t_conc = [['Código', 'Descrição', 'Comprador', 'Clientes 12m', '% do maior cliente', 'Maior cliente', 'Estoque R$',
               'Compra CLAUDE out–dez R$', 'Venda nas bases ÷ TOTVS (abr–set)']]
    for s in conc:
        t_conc.append([s['cod'], s['desc'][:60], s['comprador'], s['ncli'], s['top_share'], s['top_nome'][:40], s['vlr'],
                       s['compra3'], s['cobert']])

    agr = collections.defaultdict(lambda: [0, 0.0, 0.0])
    for s in sku:
        if s['top_share'] is not None and s['top_share'] >= 0.8:
            a = agr[s['top_nome'][:50]]
            a[0] += 1; a[1] += max(0, s['vlr']); a[2] += s['compra3']
    t_dono = [['Cliente que concentra 80%+ da venda do SKU', 'SKUs', 'Estoque R$', 'Compra CLAUDE out–dez R$']]
    for k, (n, v, c) in sorted(agr.items(), key=lambda kv: -(kv[1][1] + kv[1][2]))[:10]:
        t_dono.append([k, n, v, c])
    outros = sorted(agr.items(), key=lambda kv: -(kv[1][1] + kv[1][2]))[10:]
    t_dono.append(['Demais clientes', sum(v[0] for k, v in outros), sum(v[1] for k, v in outros), sum(v[2] for k, v in outros)])
    t_dono.append(['Total', sum(r[1] for r in t_dono[1:]), sum(r[2] for r in t_dono[1:]), sum(r[3] for r in t_dono[1:])])

    cobertos = [s for s in sku if s['ncli']]
    info = dict(data=hoje.strftime('%d/%m/%Y'), tot=tot_pos, iqr_e=estrito / tot_pos, iqr_a=amplo / tot_pos, exc=exc,
                slow=slow, phase=phase, nc=nc, livre=tot_livre, n_cli=len(cobertos), n_sku=len(sku),
                gen=t_tri[1][2], gen_n=t_tri[1][1], neg=soma(lambda s: s['vlr'] if s['vlr'] < 0 else 0),
                cli1=t_cli[1][1], cli1_rs=t_cli[1][2], cli1_compra=t_cli[1][3], conc80=t_cli[-1][5],
                gen40=sum(r[6] for r in t_gen[1:]), gen_preco=t_tri[1][4],
                conc80_rs=t_dono[-1][2], conc80_compra=t_dono[-1][3], top1=t_dono[1][1:] if len(t_dono) > 2 else None,
                c_fora=soma(lambda s: s['rs']['excesso'] + s['rs']['slow'] + s['rs']['phase'], [s for s in sku if s['abc'] == 'C']))
    br = lambda x, d=1: f'{x:,.{d}f}'.replace(',', '#').replace('.', ',').replace('#', '.')
    mi = lambda x: 'R$ ' + br(x / 1e6) + ' mi'
    pc = lambda x: br(100 * x, 0) + '%'
    abc = {r[0]: r for r in por('abc', 'Curva', ['A', 'B', 'C', ''])[1:]}
    fora = exc + slow + phase
    textos = dict(
        s7_intro=('Como foi calculado (todos os SKUs da Geral, custo do estoque = Vlr Estoque): o estoque comprometido (Reserva + VendProt) '
                  'conta como saudável; o restante é comparado ao estoque alvo CLAUDE do mês vigente (IQR estrito) e ao alvo + 2 meses de '
                  'previsão (IQR ampliado). O que passa disso é excesso, se o SKU vendeu nos últimos 6 meses (venda do TOTVS de abril a '
                  'setembro, todos os canais, ou bases de faturamento). Sem venda em 6 meses: slow moving se as bases mostram venda nos 12 '
                  'meses, senão phase out potencial.'),
        s7=[f'• {pc(info["iqr_a"])} do estoque ({mi(amplo)}) está dentro da política ampliada e {pc(info["iqr_e"])} ({mi(estrito)}) dentro do alvo '
            f'CLAUDE. Fora da política: {mi(fora)} ({pc(fora / tot_pos)}): {mi(exc)} de excesso em itens que ainda vendem, {mi(slow)} slow '
            f'moving e {mi(phase)} sem venda conhecida há 12 meses.',
            f'• A curva C tem {mi(info["c_fora"])} dos {mi(fora)} fora da política (IQR ampliado de {pc(abc["C"][4])}). As curvas A e B estão '
            f'saudáveis ({pc(abc["A"][4])} e {pc(abc["B"][4])}): o problema de estoque é a cauda, não os itens que mais vendem.',
            '• No MRP em HTML o IQR é 49% (estrito) e 59% (ampliado) porque ele só calcula política para 1.750 SKUs e deixa R$ 3,2 mi sem '
            'classificação. Aqui todos os SKUs têm alvo, porque a demanda vem da MedContr do TOTVS (todos os canais).',
            '• Phase out é hipótese, não decisão: de novembro a março só as bases de faturamento mostram a venda, e elas não têm todos os '
            'canais. Antes de liquidar, confirmar no TOTVS a última saída e se há contrato ou pedido futuro.',
            '• Por SKU: colunas "Saúde do estoque", "Última venda (mês)" e "Fora da política QTD / R$" na tabela ParamClaude (aba param '
            'CLAUDE), para filtrar por comprador.'],
        s8_intro=('Excedente livre = estoque fora da política (excesso + slow + phase out), sem o comprometido. Triagem: personalizado (nome '
                  'do cliente na descrição, PERSONALIZ/LOGO ou grupo 0015), concentrado (1 cliente com 80% ou mais da venda em 12 meses), '
                  'revisar requisitos (alimento, EPI/saúde, elétrico ou volumoso) e genérico. O preço de mercado vem das colunas Oceano, '
                  'Kalunga e Mercado Livre da Geral e só vale entre 0,3 e 5 vezes o custo (fora disso o produto comparado é outro).'),
        s8=[f'• {mi(info["gen"])} de excedente genérico em {br(info["gen_n"], 0)} SKUs são candidatos a e-commerce/marketplace; os 40 maiores '
            f'(tabela 8.1) somam {mi(info["gen40"])}. {br(info["gen_preco"], 0)} deles já têm preço de mercado pesquisado na Geral.',
            f'• Personalizados e concentrados ({mi(t_tri[3][2] + t_tri[4][2])}) não devem ir para terceiros: o caminho é o cliente dono '
            '(consumo do saldo, contrato, antecipação de pedido) ou acordo de devolução com o fornecedor.',
            '• Sugestão de piloto: pegar os 40 genéricos da tabela 8.1, conferir fisicamente (saldo, embalagem, validade), cotar o preço e '
            'os custos do canal e publicar um lote. Medir em 30 dias quanto saiu e a margem antes de ampliar.',
            '• Na triagem, o MRP em HTML achou R$ 1,9 mi de genéricos: usa outra régua de alvo e só os 1.750 SKUs com política.'],
        s9_intro=(f'Venda de outubro/2025 a setembro/2026 por cliente (raiz do CNPJ) nas bases de faturamento, sem a cópia de dezembro da '
                  f'Planilha1. {br(info["n_cli"], 0)} dos {br(info["n_sku"], 0)} SKUs têm cliente nas bases; os demais vendem por canais que '
                  'elas não trazem ou não venderam. A última coluna da tabela 9.2 mostra quanto da venda do TOTVS (abril a setembro) as '
                  'bases cobrem: perto de 1 = concentração confiável.'),
        s9=[f'• {br(info["cli1"], 0)} SKUs vendem para 1 cliente só: {mi(info["cli1_rs"])} de estoque e {mi(info["cli1_compra"])} da compra CLAUDE '
            f'de outubro a dezembro. Ao todo, {br(info["conc80"], 0)} SKUs têm 1 cliente com 80% ou mais da venda ({mi(info["conc80_rs"])} de '
            f'estoque).',
            (f'• Um único cliente (tabela 9.1) concentra {br(info["top1"][0], 0)} SKUs, {mi(info["top1"][1])} de estoque '
             f'({pc(info["top1"][1] / tot_pos)} do total) e {mi(info["top1"][2])} da compra CLAUDE do trimestre. É o maior risco de estoque '
             'parado se o cliente reduzir ou mudar o pedido.') if info['top1'] else '',
            '• Sugestão: para SKUs de 1 ou 2 clientes, comprar só com pedido, contrato ou previsão assinada pelo comercial (regime sob '
            'encomenda), com prazo de consumo e cláusula para o saldo. Revisar com o comercial, todo mês, a lista da tabela 9.2.',
            '• Na reposição normal (5 ou mais clientes) estão os SKUs em que a metodologia CLAUDE pode comprar sem consulta.'])
    out = dict(info=info, textos=textos, resumo=resumo, abc=por('abc', 'Curva', ['A', 'B', 'C', '']), comprador=por('comprador', 'Comprador'),
               grupo=por('grupo', 'Grupo'), top=t_top, triagem=t_tri, genericos=t_gen, clientes=t_cli, concentrados=t_conc, donos=t_dono)
    json.dump(out, open(S / 'saude.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=1, default=float)

    with open(S / 'param_saude.csv', 'w', encoding='utf-8-sig', newline='') as f:
        w = csv.writer(f, delimiter=';')
        w.writerow(['Codigo', 'Última venda (mês)', 'Saúde do estoque', 'Fora da política QTD', 'Fora da política R$',
                    'Triagem e-commerce', 'Clientes 12m', '% maior cliente', 'Maior cliente', 'Classe clientes'])
        for s in sku:
            w.writerow([s['cod'], s['ult'], s['classe'], round(s['livre_q'], 2), round(s['livre_rs'], 2), s['triagem'],
                        s['ncli'] or '', round(s['top_share'], 4) if s['top_share'] is not None else '', s['top_nome'], s['classe_cli']])
    print(json.dumps(info, ensure_ascii=False, indent=1, default=float))


if __name__ == '__main__':
    a = sys.argv[1:]
    main(a[0], a[1], a[2], a[3], dt.date.fromisoformat(a[4]) if len(a) > 4 else dt.date.today())
