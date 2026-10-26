"""Demanda mensal por produto a partir das bases de faturamento (fat AAAA.N.xlsx).

    python extrai_demanda.py SAIDA.pkl "fat 2023.1.xlsx" "fat 2024.1.xlsx" ... [--so-aba "fat 2025.2.xlsx:Planilha1"]

A aba Planilha1 do "fat 2025.2.xlsx" e copia exata de dezembro/2025 da Listagem: fica de fora (IGNORAR).
--so-aba ARQ:ABA le so essa aba (usado para descontar a copia de uma extracao antiga).

Le pelo nome do cabecalho (FILIAL, EMISSAO, COD_PROD, QUANT, VLR_TOTAL, COD_CLI, PEDIDO). Um processo por aba.
Saida (dict):
- mes[(filial, cod, 'AAAA-MM')] = [quantidade, valor, linhas]
- cli12[(filial, cod, cod_cli)] = quantidade dos ultimos 12 meses completos (concentracao de clientes)
- ped[(filial, cod, 'AAAA-MM')] = maior quantidade num unico pedido no mes (pico de pedido grande)
- datas min/max por arquivo
"""
import collections
import datetime as dt
import multiprocessing as mp
import os
import pickle
import sys
import warnings

import openpyxl

warnings.filterwarnings('ignore')
INICIO_12M = '2025-10'      # janela de concentracao de clientes: out/25 a set/26
IGNORAR = {('fat 2025.2.xlsx', 'Planilha1')}   # copia de dez/25 da aba Listagem


def ym(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return f'{v.year:04d}-{v.month:02d}', v.strftime('%Y-%m-%d')
    s = str(v or '').strip()
    if len(s) >= 10 and s[2] == '/':
        return f'{s[6:10]}-{s[3:5]}', f'{s[6:10]}-{s[3:5]}-{s[0:2]}'
    if len(s) == 8 and s.isdigit():
        return f'{s[0:4]}-{s[4:6]}', f'{s[0:4]}-{s[4:6]}-{s[6:8]}'
    return None, None


def num(v):
    if v is None:
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace('.', '').replace(',', '.') if ',' in str(v) else str(v).strip()
    try:
        return float(s)
    except ValueError:
        return 0.0


def trabalha(arg):
    arq, aba = arg
    wb = openpyxl.load_workbook(arq, read_only=True)
    ws = wb[aba]
    mes = collections.defaultdict(lambda: [0.0, 0.0, 0])
    cli = collections.defaultdict(float)
    pedq = collections.defaultdict(float)
    cab, dmin, dmax, n = None, '9999', '0000', 0
    for r in ws.iter_rows(values_only=True):
        if cab is None:
            if r and 'COD_PROD' in [str(c).strip() for c in r if c is not None]:
                cab = {str(c).strip(): j for j, c in enumerate(r) if c is not None}
            continue
        try:
            fil = str(r[cab['FILIAL']] or '').strip()
            cod = str(r[cab['COD_PROD']] or '').strip()
            m, d = ym(r[cab['EMISSAO']])
        except IndexError:
            continue
        if not cod or not m:
            continue
        q = num(r[cab['QUANT']])
        v = num(r[cab['VLR_TOTAL']])
        a = mes[(fil, cod, m)]
        a[0] += q
        a[1] += v
        a[2] += 1
        n += 1
        dmin, dmax = min(dmin, d), max(dmax, d)
        if m >= INICIO_12M and m <= '2026-09':
            cli[(fil, cod, str(r[cab['COD_CLI']] or '').strip())] += q
        p = str(r[cab['PEDIDO']] or '').strip() if 'PEDIDO' in cab else ''
        if p:
            pedq[(fil, cod, m, p)] += q
    wb.close()
    pk = collections.defaultdict(float)
    for (fil, cod, m, p), q in pedq.items():
        pk[(fil, cod, m)] = max(pk[(fil, cod, m)], q)
    print(arq, aba, n, dmin, dmax, file=sys.stderr, flush=True)
    return dict(mes), dict(cli), dict(pk), (arq, aba, n, dmin, dmax)


def main(saida, arquivos, so_aba=None):
    tarefas = []
    for a in arquivos:
        wb = openpyxl.load_workbook(a, read_only=True)
        tarefas += [(a, s) for s in wb.sheetnames if (os.path.basename(a), s) not in IGNORAR]
        wb.close()
    if so_aba:
        tarefas = [tuple(so_aba.split(':', 1))]
    with mp.Pool(4) as pool:
        res = pool.map(trabalha, tarefas)
    mes = collections.defaultdict(lambda: [0.0, 0.0, 0])
    cli = collections.defaultdict(float)
    pk = collections.defaultdict(float)
    info = []
    for m_, c_, p_, i_ in res:
        for k, v in m_.items():
            a = mes[k]
            a[0] += v[0]; a[1] += v[1]; a[2] += v[2]
        for k, v in c_.items():
            cli[k] += v
        for k, v in p_.items():
            pk[k] = max(pk[k], v)
        info.append(i_)
    pickle.dump({'mes': dict(mes), 'cli12': dict(cli), 'ped': dict(pk), 'info': info}, open(saida, 'wb'))
    print('ok', len(mes), 'produto-mes', file=sys.stderr)


if __name__ == '__main__':
    a = sys.argv[1:]
    so = a[a.index('--so-aba') + 1] if '--so-aba' in a else None
    arqs = [x for x in a[1:] if x != '--so-aba' and x != so]
    main(a[0], arqs, so)
