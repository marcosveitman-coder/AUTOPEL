"""Acha os SKUs em que o Disponivel (coluna O) veio do TOTVS sem a virgula decimal.

    python diagnostico_disponivel.py "Relatorio Revenda atualizado.xlsx"

Le as abas Geral (arq. ...) e a copia de ontem. Para cada SKU refaz a conta do TOTVS,
Comprar = PontoPedido - (Disponivel - VendProt + Comprado), e mostra os que nao batem com:
- o Disponivel verdadeiro (PontoPedido - Comprar + VendProt - Comprado);
- quantas vezes o gravado e maior (10, 100, 10000...: digitos colados);
- o EstDias gravado e o certo.
"""
import sys

import openpyxl


def num(x):
    try:
        return float(x)
    except (TypeError, ValueError):
        return 0.0


def main(arq):
    wb = openpyxl.load_workbook(arq, read_only=True, data_only=True)
    for aba in [s for s in wb.sheetnames if s.startswith('Geral (arq')]:
        R = list(wb[aba].iter_rows(values_only=True))
        lc = next(i for i, r in enumerate(R[:6]) if r and r[0] == 'Comprador')
        cab = {str(c).strip(): j for j, c in enumerate(R[lc]) if c is not None}
        ruins, frac_o, n = [], 0, 0
        for i, r in enumerate(R[lc + 1:], start=lc + 2):
            if not r or r[cab['Codigo']] in (None, ''):
                continue
            n += 1
            E, F, O, P, C, I, W = (num(r[cab[k]]) for k in
                                   ('PontoPedido', 'Comprar.', 'Disponivel', 'VendProt', 'Comprado', 'MedContr', 'EstDias'))
            real = E - F + P - C
            if abs(real - round(real)) > 1e-9:
                frac_o += 1
            if abs(F - (E - (O - P + C))) > 0.01:
                certo = (real - P) / I * 30 if I else 0
                ruins.append(f'  linha {i} {r[cab["Codigo"]]}: Disponivel gravado {O:g}, verdadeiro {real:g} '
                             f'(x{O / real:g}); EstDias gravado {W:,.1f}, certo {certo:,.1f}')
        print(f'[{aba}] {n} SKUs; Disponivel com decimal: {frac_o}; conta nao bate: {len(ruins)}')
        print('\n'.join(ruins))
    wb.close()


if __name__ == '__main__':
    main(sys.argv[1])
