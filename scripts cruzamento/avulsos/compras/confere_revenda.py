"""Confere a revisao do Relatorio Revenda: compara o arquivo revisado com o original celula a celula.

    python confere_revenda.py ORIGINAL.xlsx REVISADO.xlsx

- aba Geral (a da Table1): colunas A a X (TOTVS) tem que ficar iguais; lista o que mudou nas outras colunas;
  refaz em Python a programacao do trimestre e compara com o que o Excel calculou.
- Demais abas: conta as celulas que mudaram (valor ou formula), por aba e coluna.
"""
import collections
import sys

import openpyxl
from openpyxl.utils import get_column_letter as L



def ler(arq):
    out = {}
    for so_valor in (False, True):
        wb = openpyxl.load_workbook(arq, read_only=True, data_only=so_valor)
        for ws in wb.worksheets:
            out.setdefault(ws.title, {})['v' if so_valor else 'f'] = [r for r in ws.iter_rows(values_only=True)]
        wb.close()
    return out


def cel(rows, i, j):
    return rows[i][j] if i < len(rows) and j < len(rows[i]) else None


def num(x):
    try:
        return float(x)
    except (TypeError, ValueError):
        return 0.0


def igual(a, b):
    # formula matricial vem como objeto (ArrayFormula), sem comparacao propria: compara o texto
    a, b = getattr(a, 'text', a), getattr(b, 'text', b)
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return abs(a - b) < 1e-6
    return a == b


def compara_aba(a, b, ncols=200):
    dif = collections.Counter()
    for k in ('f', 'v'):
        A, B = a.get(k, []), b.get(k, [])
        for i in range(max(len(A), len(B))):
            for j in range(min(ncols, max(len(A[i]) if i < len(A) else 0, len(B[i]) if i < len(B) else 0))):
                if not igual(cel(A, i, j), cel(B, i, j)):
                    dif[(k, L(j + 1))] += 1
    return dif


def main(orig, rev):
    O, N = ler(orig), ler(rev)
    # aba Geral: comeca com 'Geral (arq'; a copia oculta de ontem tem o mesmo comeco e nome mais comprido
    GERAL = min((s for s in O if s.startswith('Geral (arq')), key=len)
    print('abas novas:', [s for s in N if s not in O], '| abas sumidas:', [s for s in O if s not in N])

    # 1) Geral A:X
    go, gn = O[GERAL], N[GERAL]
    ruins = 0
    for k in ('f', 'v'):
        for i in range(max(len(go[k]), len(gn[k]))):
            for j in range(24):
                if not igual(cel(go[k], i, j), cel(gn[k], i, j)):
                    ruins += 1
    print(f'Geral A:X celulas diferentes: {ruins}')
    d = compara_aba(go, gn)
    print('Geral, colunas que mudaram:', sorted(d.items(), key=lambda x: (x[0][0], len(x[0][1]), x[0][1])))

    # 2) programacao refeita em Python
    V = gn['v']
    cab = {str(c): j for j, c in enumerate(V[1]) if c is not None}
    # colunas de previsao: nome novo (desde 23/10/2026) ou o antigo
    antigos = {'Compra mês vigente QTD': 'GPT previsao OUTUBRO QNTDD', 'Compra próximo mês QTD': 'GPT previsao NOVEMBRO QNTDD',
               'Compra mês subsequente QTD': 'GPT  previsao DEZEMVRO QNTDD', 'Compra mês vigente R$': 'GPT  previsao OUTUBRO R$',
               'Compra próximo mês R$': 'GPT  previsao NOVEMBRO R$', 'Compra mês subsequente R$': 'GPT previsao DEZEMVRO R$'}
    for novo, velho in antigos.items():
        if novo not in cab and velho in cab:
            cab[novo] = cab[velho]

    def c(i, nome):
        return V[i][cab[nome]]
    erros = collections.Counter()
    for i in range(2, len(V)):
        if c(i, 'Codigo      ') in (None, ''):
            continue
        E, F, G, I = (num(c(i, n)) for n in (' PontoPedido', 'Comprar.', ' Ult.Preco', ' MedContr'))
        s0 = E - F
        m0 = max(0, -s0) if E == 0 else max(0, E - s0)
        m1 = 0 if E == 0 else max(0, E - (s0 + m0 - I))
        m2 = 0 if E == 0 else max(0, E - (s0 + m0 + m1 - 2 * I))
        esperado = {
            'Saldo hoje (Disp-VendProt+Comprado)': s0,
            'Compra mês vigente QTD': m0, 'Compra próximo mês QTD': m1, 'Compra mês subsequente QTD': m2,
            'Compra mês vigente R$': m0 * G, 'Compra próximo mês R$': m1 * G, 'Compra mês subsequente R$': m2 * G,
            'Compra trimestre QTD': m0 + m1 + m2, 'Compra trimestre R$': (m0 + m1 + m2) * G,
        }
        for n, e in esperado.items():
            if abs(num(c(i, n)) - e) > 0.01:
                erros[n] += 1
        q, comprar = num(c(i, ' Ruptura')), m0
        sit = ('1 RUPTURA' if q < 0 else '2 COMPRAR AGORA' if comprar > 0 else '3 COMPRAR NOS PROXIMOS MESES'
               if m0 + m1 + m2 > 0 else '5 SEM PONTO DE PEDIDO' if E == 0 else '4 ESTOQUE OK')
        if c(i, 'Situacao compra') != sit:
            erros['Situacao compra'] += 1
        div = abs(F - (E - (num(c(i, ' Disponivel')) - num(c(i, ' VendProt')) + num(c(i, ' Comprado'))))) > 0.5
        if (c(i, 'divergencia COMPRAR') == 'divergente') != div:
            erros['divergencia COMPRAR'] += 1
    print('programacao refeita em Python, linhas que nao batem:', dict(erros) or 'nenhuma')
    tot ={n: V[0][cab[n]] for n in cab if cab[n] < len(V[0]) and isinstance(V[0][cab[n]], (int, float))}
    print('totais da linha 1:', {k: round(v) for k, v in tot.items()})

    # 3) demais abas
    for s in O:
        if s == GERAL or s not in N:
            continue
        d = compara_aba(O[s], N[s])
        if d:
            print(f'[{s}] mudou:', sorted(d.items(), key=lambda x: (x[0][0], len(x[0][1]), x[0][1])))
    return N


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
