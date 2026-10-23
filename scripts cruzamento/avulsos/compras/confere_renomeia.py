"""Confere a troca de nomes das colunas de previsao: o layout tem que ficar igual.

    python confere_renomeia.py ANTES.xlsx DEPOIS.xlsx

- valores de todas as abas (so podem mudar cabecalhos e titulos);
- estilo de cada celula, larguras de coluna e alturas de linha (pelo XML de cada aba);
- formatacoes proprias de cada tabela dinamica (quantidade e conteudo do bloco <formats>) e a posicao dos campos.
"""
import collections
import re
import sys
import zipfile

import openpyxl
from openpyxl.utils import get_column_letter as L


def abas_xml(z):
    wb = z.read('xl/workbook.xml').decode('utf-8')
    rel = z.read('xl/_rels/workbook.xml.rels').decode('utf-8')
    rm = {re.search(r'Id="([^"]+)"', e).group(1): re.search(r'Target="([^"]+)"', e).group(1)
          for e in re.findall(r'<Relationship [^>]+>', rel)}
    return {n: 'xl/' + rm[r].replace('/xl/', '') for n, r in
            re.findall(r'<sheet name="([^"]+)" sheetId="\d+"(?: state="\w+")? r:id="(rId\d+)"', wb)}


def estilos(z):
    """Assinatura de cada estilo de celula (cellXfs) e de cada formato diferencial (dxfs), pelo conteudo:
    o Excel renumera os estilos ao salvar, entao o indice sozinho nao serve para comparar."""
    x = z.read('xl/styles.xml').decode('utf-8')
    def bloco(tag):
        m = re.search(r'<' + tag + r'\b[^>]*?(?:/>|>(.*?)</' + tag + '>)', x, re.S)
        return (m.group(1) or '') if m else ''

    def itens(tag, filho):
        return re.findall(r'<' + filho + r'\b(?:[^>]*/>|[^>]*>.*?</' + filho + '>)', bloco(tag), re.S)

    fmts = dict(re.findall(r'<numFmt numFmtId="(\d+)" formatCode="([^"]*)"', bloco('numFmts')))
    # <family> e so dica de familia da fonte (nao muda a aparencia); o Excel junta fontes que so diferem nela
    fonts = [re.sub(r'<family val="\d+"/>', '', f) for f in itens('fonts', 'font')]
    fills, borders = itens('fills', 'fill'), itens('borders', 'border')
    sig = []
    for xf in itens('cellXfs', 'xf'):
        a = dict(re.findall(r'(\w+)="([^"]*)"', xf.split('>')[0]))
        nf = a.get('numFmtId', '0')
        sig.append((fmts.get(nf, nf), fonts[int(a.get('fontId', 0))], fills[int(a.get('fillId', 0))],
                    borders[int(a.get('borderId', 0))], re.sub(r'\s+', ' ', xf.split('>', 1)[1] if '>' in xf else '')))
    dxfs = itens('dxfs', 'dxf')
    return sig, dxfs


def layout(x, sig):
    estilo = {r: sig[int(s or 0)] for r, s in re.findall(r'<c r="([A-Z]+\d+)"(?: s="(\d+)")?', x)}
    # largura de cada coluna (o estilo da coluna e o indicador de melhor ajuste nao mudam a largura)
    cols = [re.sub(r' (?:style|bestFit)="\d+"', '', c) for c in re.findall(r'<col [^>]*/>', x)]
    rows = dict(re.findall(r'<row r="(\d+)"[^>]*? ht="([\d.]+)"', x))
    return estilo, cols, rows


def dinamicas(z, dxfs):
    out = {}
    for x in [n for n in z.namelist() if re.match(r'xl/pivotTables/pivotTable\d+\.xml$', n)]:
        t = z.read(x).decode('utf-8')
        nome = re.search(r' name="([^"]+)"', t).group(1)
        loc = re.search(r'<location ref="([^"]+)"', t).group(1)
        fm = re.search(r'<formats count="\d+">.*?</formats>', t, re.S)
        fm = re.sub(r'dxfId="(\d+)"', lambda m: 'dxf=' + dxfs[int(m.group(1))], fm.group(0)) if fm else None
        campos = re.findall(r'<dataField name="[^"]*" fld="(\d+)"', t)
        out[(nome, loc)] = (fm or '', campos)
    return out


def main(antes, depois):
    za, zd = zipfile.ZipFile(antes), zipfile.ZipFile(depois)
    xa, xd = abas_xml(za), abas_xml(zd)
    assert list(xa) == list(xd), 'abas diferentes'
    sa, dxa = estilos(za)
    sd, dxd = estilos(zd)

    ruins = 0
    for aba in xa:
        ea, ca, ra = layout(za.read(xa[aba]).decode('utf-8'), sa)
        ed, cd, rd = layout(zd.read(xd[aba]).decode('utf-8'), sd)
        de = sum(1 for k in set(ea) | set(ed) if ea.get(k) != ed.get(k))
        if de or ca != cd or ra != rd:
            ruins += 1
            alt = sorted((k for k in set(ra) | set(rd) if ra.get(k) != rd.get(k)), key=int)
            print(f'[{aba}] estilos diferentes: {de}; colunas iguais: {ca == cd}; linhas com altura diferente: {alt[:12]}')
            if ca != cd:
                print('    colunas antes :', [c for c in ca if c not in cd][:6])
                print('    colunas depois:', [c for c in cd if c not in ca][:6])
    print('abas com layout diferente:', ruins)

    pa, pd = dinamicas(za, dxa), dinamicas(zd, dxd)
    print('dinamicas:', len(pa), '->', len(pd))
    for k in pa:
        if k not in pd:
            print('  sumiu/mudou de lugar:', k)
        elif pa[k] != pd[k]:
            print('  formatacao ou campos diferentes:', k, pa[k][1], pd[k][1])

    va = openpyxl.load_workbook(antes, read_only=True, data_only=True)
    vd = openpyxl.load_workbook(depois, read_only=True, data_only=True)
    for aba in va.sheetnames:
        A = list(va[aba].iter_rows(values_only=True))
        D = list(vd[aba].iter_rows(values_only=True))
        dif = []
        for i in range(max(len(A), len(D))):
            ra_ = A[i] if i < len(A) else ()
            rd_ = D[i] if i < len(D) else ()
            for j in range(min(200, max(len(ra_), len(rd_)))):
                a = ra_[j] if j < len(ra_) else None
                d = rd_[j] if j < len(rd_) else None
                if a != d and not (isinstance(a, float) and isinstance(d, float) and abs(a - d) < 1e-6):
                    dif.append(f'{L(j + 1)}{i + 1}: {a!r} -> {d!r}')
        if dif:
            print(f'[{aba}] {len(dif)} valores mudaram:')
            for t in dif[:14]:
                print('   ', t)
    va.close()
    vd.close()


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
