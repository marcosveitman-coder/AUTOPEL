"""Grava a base consolidada na aba "FARMACIA CONSOLIDADO" (sheet18.xml) do MAPA MERCADO, so trocando o XML dessa aba."""
import os, re, sys, pickle, zipfile, datetime, collections
from xml.sax.saxutils import escape
from openpyxl.utils import get_column_letter

MAPA = os.environ.get("AUTOPEL_MAPA", "MAPA MERCADO.xlsx")
SHEET = "xl/worksheets/sheet18.xml"
S_HEAD, S_TXT, S_NUM, S_DATE = 200, 0, 127, 107
CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")

COLS = [("REDE / ASSOCIAÇÃO", 34), ("TIPO", 24), ("NOME FANTASIA", 40), ("RAZÃO SOCIAL", 48), ("CNPJ", 20),
        ("Nº DE LOJAS (RECEITA)", 12), ("Nº DE LOJAS (PESQUISA)", 12), ("UF SEDE", 7), ("MUNICÍPIO SEDE", 22),
        ("SITUAÇÃO NA RECEITA", 12), ("NA ABA FARMACIA?", 10), ("VALIDAÇÃO / OBSERVAÇÃO", 60),
        ("DATA ÚLTIMA COMPRA", 13), ("STATUS", 13), ("COMO FOI ACHADO NO FATURAMENTO", 28), ("CLIENTE ENCONTRADO NO FATURAMENTO", 60)]


def fmt(c):
    c = re.sub(r"\D", "", c or "")
    return f"{c[:2]}.{c[2:5]}.{c[5:8]}/{c[8:12]}-{c[12:]}" if len(c) == 14 else c


def excel_date(dt):
    d = datetime.date(dt // 10000, dt // 100 % 100, dt % 100)
    return (d - datetime.date(1899, 12, 30)).days


def cell(ref, v, s):
    if v is None or v == "":
        return ""
    if isinstance(v, (int, float)):
        return f'<c r="{ref}" s="{s}"><v>{v}</v></c>'
    t = escape(CTRL.sub("", str(v)))
    return f'<c r="{ref}" s="{s}" t="inlineStr"><is><t xml:space="preserve">{t}</t></is></c>'


def main():
    B = pickle.load(open("cruz.pkl", "rb")); rows = B["rows"]
    MUN = pickle.load(open("municipios.pkl", "rb"))
    for x in rows:
        if x.get("mun") and x["mun"].isdigit():
            x["mun"] = MUN.get(x["mun"], x["mun"])
    # rede-mae (Farmarcas): herda a compra mais recente das suas bandeiras
    for mae in [x for x in rows if x["tipo"].startswith("Associação") and not x.get("mae")]:
        filhas = [x for x in rows if x.get("mae") == mae["fant"] and x.get("ult")]
        if filhas:
            f = max(filhas, key=lambda z: z["ult"]["dt"])
            if not mae.get("ult") or f["ult"]["dt"] > mae["ult"]["dt"]:
                mae["ult"] = dict(f["ult"], how=f'Bandeira {f["fant"]}: ' + f["ult"]["how"])
                d = f["ult"]["dt"]; dd = datetime.date(d // 10000, d // 100 % 100, d % 100)
                mae["status"] = "ATIVO" if (datetime.date(2026, 10, 21) - dd).days <= 360 else "A TRABALHAR"
    ordem = {"Associação / franquia": 0, "Associação / franquia (CNPJ próprio)": 0, "Rede (CNPJ próprio)": 1, "Farmácia": 2}
    rows.sort(key=lambda x: (ordem.get(x["tipo"], 3), -(x["lojas"] or 0), x["uf"] or "", x["fant"] or ""))
    out = ['<row r="1" spans="1:16" ht="45" customHeight="1">' +
           "".join(cell(f"{get_column_letter(j+1)}1", h, S_HEAD) for j, (h, w) in enumerate(COLS)) + "</row>"]
    n = 1
    for x in rows:
        n += 1
        u = x.get("ult")
        cli = ""
        if u:
            cli = f'{u.get("cliente", "")} | CNPJ {fmt(u.get("cnpj", ""))}'
            if u.get("associada"): cli += f' | associada: {u["associada"]}'
            if u.get("produto"): cli += f' | produto: {u["produto"]}'
            if u.get("dominio"): cli += f' | e-mail @{u["dominio"]}'
        vals = [(x["rede"], S_TXT), (x["tipo"], S_TXT), (x["fant"], S_TXT), (x["rz"], S_TXT), (fmt(x["cnpj"]), S_TXT),
                (x["lojas"] if x["lojas"] is not None else None, S_NUM), (x["inf"] if isinstance(x["inf"], (int, float)) else None, S_NUM),
                (x["uf"], S_TXT), (x["mun"], S_TXT), (x["sit"], S_TXT), (x["aba"], S_TXT), (x["obs"], S_TXT),
                (excel_date(u["dt"]) if u else None, S_DATE), (x["status"], S_TXT), (u["how"] if u else "", S_TXT), (cli, S_TXT)]
        out.append(f'<row r="{n}" spans="1:16">' + "".join(cell(f"{get_column_letter(j+1)}{n}", v, s) for j, (v, s) in enumerate(vals)) + "</row>")
    last = get_column_letter(len(COLS))
    cols = "<cols>" + "".join(f'<col min="{j+1}" max="{j+1}" width="{w}" customWidth="1"/>' for j, (h, w) in enumerate(COLS)) + "</cols>"
    z = zipfile.ZipFile(MAPA)
    x = z.read(SHEET).decode("utf8")
    ja = re.search(r"<cols>.*?</cols><sheetData>.*?</sheetData><autoFilter[^>]*/>", x, re.S)
    if ja and "REDE / ASSOCIA" in x[:20000]:  # aba ja gravada por este script: refaz
        x = x.replace(ja.group(0), "<sheetData/>")
        x = re.sub(r'<dimension ref="[^"]*"/>', '<dimension ref="A1"/>', x)
        x = re.sub(r"<sheetViews>.*?</sheetViews>", '<sheetViews><sheetView workbookViewId="0"/></sheetViews>', x, flags=re.S)
    assert "<sheetData/>" in x, "a aba FARMACIA CONSOLIDADO tem outro conteudo - nao sobrescrevo"
    x = x.replace('<dimension ref="A1"/>', f'<dimension ref="A1:{last}{n}"/>')
    x = x.replace('<sheetViews><sheetView workbookViewId="0"/></sheetViews>',
                  '<sheetViews><sheetView workbookViewId="0"><pane ySplit="1" topLeftCell="A2" activePane="bottomLeft" state="frozen"/>'
                  '<selection pane="bottomLeft" activeCell="A2" sqref="A2"/></sheetView></sheetViews>')
    x = x.replace("<sheetData/>", cols + "<sheetData>" + "".join(out) + "</sheetData>" + f'<autoFilter ref="A1:{last}{n}"/>')
    tmp = MAPA + ".tmp"
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zo:
        for info in z.infolist():
            zo.writestr(info, x.encode("utf8") if info.filename == SHEET else z.read(info.filename), compress_type=zipfile.ZIP_DEFLATED)
    z.close()
    os.replace(tmp, MAPA)
    print("gravado:", n - 1, "linhas")


if __name__ == "__main__":
    try:
        open(MAPA, "r+b").close()
    except PermissionError:
        raise SystemExit("BLOQUEADO - feche o MAPA MERCADO no Excel")
    main()
