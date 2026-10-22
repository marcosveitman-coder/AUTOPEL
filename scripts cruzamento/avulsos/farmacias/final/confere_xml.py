import zipfile, sys, openpyxl, pickle, datetime as dt
from openpyxl.utils import get_column_letter as L
a,b=sys.argv[1],sys.argv[2]
za,zb=zipfile.ZipFile(a),zipfile.ZipFile(b)
na,nb=set(za.namelist()),set(zb.namelist())
print('partes só num lado:', na^nb)
print('partes diferentes:', [n for n in sorted(na&nb) if za.read(n)!=zb.read(n)])
wa=openpyxl.load_workbook(a,read_only=True,data_only=True); wb=openpyxl.load_workbook(b,read_only=True,data_only=True)
perm={'"FARMACIA CONSOLIDADO"':set('NOPQRST'),'"FARMACIA VF"':set('OPQRVW')}
R=pickle.load(open('farm_final.pkl','rb'))
for nome,key,cd,cst in (('"FARMACIA CONSOLIDADO"','co','N','O'),('"FARMACIA VF"','vf','O','P')):
    A=list(wa[nome].iter_rows(values_only=True)); B=list(wb[nome].iter_rows(values_only=True))
    fora=0; mud=__import__('collections').Counter()
    for i in range(max(len(A),len(B))):
        ra=A[i] if i<len(A) else (); rb=B[i] if i<len(B) else ()
        for j in range(max(len(ra),len(rb))):
            va=ra[j] if j<len(ra) else None; vb=rb[j] if j<len(rb) else None
            if va!=vb:
                c=L(j+1)
                if c in perm[nome]: mud[c]+=1
                else: fora+=1; print('FORA', nome, c, i+1, va, vb) if fora<5 else None
    print(nome, 'linhas', len(A), len(B), 'mudanças fora das colunas pedidas:', fora, 'por coluna:', dict(mud))
    X=R[key]; hdr=B[0]
    # confere gravado == calculado
    ci={h:j for j,h in enumerate(hdr)}
    errs=0
    for i in range(X['n']):
        r=B[i+1]
        if r[ci['VENDEDOR']] != (X['vend'][i] or None): errs+=1
        if r[ci['Total de caixas/ mês']] != int(X['cx'][i]): errs+=1
        d=X['best'][i]['dt']; dv=r[ci['DATA ÚLTIMA COMPRA']]
        if (d.date() if d else None)!=(dv.date() if isinstance(dv,dt.datetime) else None): errs+=1
        if r[ci['STATUS']]!=X['status'][i]: errs+=1
    print('   cabeçalhos novos:', hdr[-3:], ' divergências gravado×calculado:', errs)
wa.close(); wb.close()
