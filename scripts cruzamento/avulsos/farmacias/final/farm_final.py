"""Fechamento das abas "FARMACIA CONSOLIDADO" e "FARMACIA VF" do MAPA MERCADO (22/10/2026).

Para cada linha:
  1. ultima compra = a mais recente entre a data que ja esta na aba e as fontes novas do relatorio
     "ultimo fat" do TOTVS (DT_FAT por cliente, desde 2019):
       - raiz do CNPJ da propria linha;
       - raiz do comprador ja identificado (coluna "cliente encontrado");
       - grupo do cadastro (coluna Descricao): so para a linha-cabeca de rede/associacao, pelo mapa grupos_redes.json;
       - linhas da mesma rede (associadas/outras raizes) sobem para a cabeca; as bandeiras sobem para a Farmarcas.
     STATUS = ATIVO se a ultima compra tem ate 360 dias.
  2. vendedor da carteira (CLIENTES): comprador da ultima compra (CNPJ, depois raiz), depois a propria linha;
     carteiras que nao sao pessoas ficam de fora; recurso: vendedor da ultima venda (faturamento).
  3. porte A-D: o mesmo da FARMACIA VF (coluna U) para a mesma empresa; sem par, a mesma regra (lojas x R$ 2,2 mi).
  4. total de caixas/mes = caixas por loja do porte (A 17, B 12, C 7, D 4) x lojas.  Nenhuma loja e contada duas vezes:
     a base e o universo da Receita, entao a linha de associacao/franquia (e a Farmarcas) fica com 0 ou so com as lojas
     da propria raiz; a linha da rede com CNPJ proprio fica com o nº considerado menos as outras raizes com linha.

Uso:  python farm_final.py            -> so confere e grava farm_final.pkl (nada e gravado na planilha)
      python farm_final.py --gravar   -> grava as duas abas pelo XML (feche o arquivo no Excel antes)
Variaveis: AUTOPEL_MAPA, AUTOPEL_ULTFAT (ultimo fat .xlsx), AUTOPEL_CLIENTES_XLSX, AUTOPEL_AGG (pasta dos .pkl do
faturamento), AUTOPEL_GRUPOS (grupos_redes.json, local e fora do git: nomes de grupos de clientes).
"""
import os, re, sys, json, glob, pickle, datetime as dt, collections as C, unicodedata
import openpyxl

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", ".."))
from gravar_coluna_xml import gravar_celulas

MAPA = os.environ.get("AUTOPEL_MAPA", "MAPA MERCADO.xlsx")
ULTFAT = os.environ.get("AUTOPEL_ULTFAT", "ultimo fat.xlsx")
CLIENTES = os.environ.get("AUTOPEL_CLIENTES_XLSX", "CLIENTES.xlsx")
AGG = os.environ.get("AUTOPEL_AGG", os.path.join(os.environ.get("TEMP", "."), "cruzamento_supermercados", "agg"))
GRUPOS = os.environ.get("AUTOPEL_GRUPOS", os.path.join(os.path.dirname(os.path.abspath(__file__)), "grupos_redes.json"))
HOJE = dt.datetime(2026, 10, 22)
INICIO = dt.datetime(2019, 1, 1)
CX = {"A": 17, "B": 12, "C": 7, "D": 4}
POR_LOJA = 2.2e6
FAIXAS = [(1e9, "A – Grande (≥ R$ 1 bi)"), (3e8, "B – Média (R$ 300 mi a 1 bi)"),
          (2e7, "C – Regional (R$ 20 mi a 300 mi)"), (0, "D – Pequena / local (< R$ 20 mi)")]
NAO_PESSOA = {"INATIVOS", "CLIENTE SEM COMPRA", "RD", "REPRESENTANTE", "AUTOPEL - VENDA ENCERRADA", ""}
BANDEIRAS_FARMARCAS = "(FARMARCAS)"

D = lambda x: re.sub(r"\D", "", str(x or ""))
s = lambda x: str(x).strip() if x is not None else ""
sa = lambda x: unicodedata.normalize("NFKD", str(x or "")).encode("ascii", "ignore").decode().upper()
fmt = lambda c: f"{c[:2]}.{c[2:5]}.{c[5:8]}/{c[8:12]}-{c[12:]}" if len(c) == 14 else c


def data_br(t):
    try:
        return dt.datetime.strptime(s(t), "%d/%m/%Y")
    except ValueError:
        return None


# ---------------------------------------------------------------- fontes
def ler_listagem(f):
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    wb.close()
    return rows[2:]  # linha 1 = titulo do relatorio, linha 2 = cabecalho


def dv_ok(c):
    if len(c) != 14 or len(set(c)) == 1:
        return False
    def dig(base, pesos):
        r = sum(int(a) * b for a, b in zip(base, pesos)) % 11
        return "0" if r < 2 else str(11 - r)
    p1 = [5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2]
    return c[12] == dig(c[:12], p1) and c[13] == dig(c[:13], [6] + p1)


def fontes():
    fat_root = {}
    for f in glob.glob(os.path.join(AGG, "*.pkl")):
        if "2018" in os.path.basename(f):
            continue
        for k, v in pickle.load(open(f, "rb"))["root"].items():
            o = fat_root.get(k)
            if o is None or v[0] > o[0]:
                fat_root[k] = [v[0], dict(v[1])]
            elif v[0] == o[0]:
                for n, val in v[1].items():
                    o[1][n] = o[1].get(n, 0) + val
    # o "ultimo fat" so vale depois do fim das bases de faturamento: antes disso a data e a das bases (as poucas notas
    # anteriores que so aparecem no "ultimo fat" nao estao nas bases de venda e nao entram)
    fim = dt.datetime.strptime(str(max(v[0] for v in fat_root.values())), "%Y%m%d")
    uf = ler_listagem(ULTFAT)
    cli = ler_listagem(CLIENTES)
    vname, clientes, root_ult = {}, [], {}
    for r in uf:
        if s(r[14]) and s(r[15]):
            vname.setdefault(s(r[14]), s(r[15]))
        c = D(r[2])
        d = data_br(r[24])
        x = dict(cod=s(r[0]), loja=s(r[1]), cnpj=c, grp=s(r[3]), gdesc=s(r[4]), nome=s(r[5]), fant=s(r[6]),
                 seg=s(r[13]), dt=d if d and d > fim else None)
        clientes.append(x)
        if len(c) == 14 and x["dt"] and (c[:8] not in root_ult or x["dt"] > root_ult[c[:8]]["dt"]):
            root_ult[c[:8]] = x
    cad_cnpj, cad_root = {}, C.defaultdict(C.Counter)
    for r in cli:
        c, v = D(r[4]), s(r[5])
        if len(c) != 14 or not v:
            continue
        cad_cnpj.setdefault(c, v)
        cad_root[c[:8]][v] += 1
    return vname, clientes, root_ult, cad_cnpj, cad_root, fat_root, fim


# ---------------------------------------------------------------- abas
LAYOUT = {  # indices (0-based) das colunas de cada aba
    '"FARMACIA CONSOLIDADO"': dict(rede=0, tipo=1, fant=2, cnpj=5, lr=6, lp=7, mun=9, dt=13, st=14, how=15, cli=16,
                                   col_dt="N", col_st="O", col_how="P", col_cli="Q",
                                   novas={"vend": "R", "porte": "S", "cx": "T"}, cab=194),
    '"FARMACIA VF"': dict(rede=1, tipo=2, fant=4, cnpj=6, lr=8, lp=9, lc=7, mun=11, dt=14, st=15, how=16, cli=17, porte=20,
                          col_dt="O", col_st="P", col_how="Q", col_cli="R",
                          novas={"vend": "V", "cx": "W"}, cab=199),
}


def ler_abas():
    wb = openpyxl.load_workbook(MAPA, read_only=True, data_only=True)
    out = {n: list(wb[n].iter_rows(values_only=True)) for n in LAYOUT}
    wb.close()  # read_only segura o arquivo; sem fechar, o Windows nega a troca do arquivo
    return out


def porte_por(fat, lojas):
    v = fat if isinstance(fat, (int, float)) and fat > 0 else (lojas or 0) * POR_LOJA if isinstance(lojas, (int, float)) else 0
    return next(t for lim, t in FAIXAS if v >= lim)


def processa(nome, rows, F, grupos, vf=None):
    vname, clientes, root_ult, cad_cnpj, cad_root, fat_root, fim = F
    L = LAYOUT[nome]
    hdr, rows = rows[0], rows[1:]
    n = len(rows)
    while n and not any(v not in (None, "") for v in rows[n - 1]):
        n -= 1
    rows = rows[:n]
    g = lambda i, k: rows[i][L[k]] if len(rows[i]) > L[k] else None
    tipo = [s(g(i, "tipo")) for i in range(n)]
    rede = [s(g(i, "rede")) for i in range(n)]
    cnpj = [D(g(i, "cnpj")) for i in range(n)]

    # cabeca de cada rede: associacao; senao a linha de rede com nº de lojas da pesquisa/ranking
    heads = {}
    for i in range(n):
        if rede[i] and tipo[i].startswith("Associação") and rede[i] not in heads:
            heads[rede[i]] = i
    for i in range(n):
        if rede[i] and rede[i] not in heads and tipo[i].startswith("Rede") and isinstance(g(i, "lp"), (int, float)):
            heads[rede[i]] = i
    head_set = set(heads.values())

    def partes(lab):
        if lab in heads:
            return [lab]
        return [p for p in lab.split(" / ") if p in heads]

    # ---- 1. ultima compra
    best = []
    for i in range(n):
        d0 = g(i, "dt")
        best.append(dict(dt=d0 if isinstance(d0, dt.datetime) else None, how=s(g(i, "how")), cli=s(g(i, "cli")),
                         fonte="aba", comprador=""))
    m = re.search
    homonimos = 0
    for i in range(n):
        b = best[i]
        q = m(r"CNPJ (\d{2}\.\d{3}\.\d{3}/\d{4}-\d{2})", b["cli"])
        b["comprador"] = D(q.group(1)) if q else ""
        # razao social igual mas a linha tem CNPJ valido achado na Receita e o comprador e de outra raiz: homonimo
        if (b["how"].startswith("Razão social") and dv_ok(cnpj[i]) and s(g(i, "mun"))
                and b["comprador"][:8] != cnpj[i][:8]):
            b.update(dt=None, how="", cli="", comprador="", fonte="homônimo removido")
            homonimos += 1
        cand = []
        if len(cnpj[i]) == 14 and cnpj[i][:8] in root_ult:
            cand.append(("raiz", root_ult[cnpj[i][:8]]))
        if len(b["comprador"]) == 14 and b["comprador"][:8] in root_ult:
            cand.append(("comprador", root_ult[b["comprador"][:8]]))
        for how, x in cand:
            if b["dt"] is None or x["dt"] > b["dt"]:
                b.update(dt=x["dt"], fonte="ultimo fat (" + how + ")", comprador=x["cnpj"],
                         cli=f"{x['nome']} | CNPJ {fmt(x['cnpj'])}")
                if how == "raiz":
                    b["how"] = "CNPJ raiz"
                elif not b["how"]:
                    b["how"] = "CNPJ raiz do comprador"

    # grupos do cadastro -> cabeca da rede
    por_grupo = C.defaultdict(list)
    for x in clientes:
        if x["grp"] in grupos["grupos"] and (x["seg"] == "DROGARIAS" or re.search(r"FARM|DROG|PHARM", sa(x["nome"] + " " + x["fant"]))):
            por_grupo[x["grp"]].append(x)

    def dt_cli(x):  # ultima venda da raiz do cliente: bases de faturamento (2019+) e, depois delas, o "ultimo fat"
        if len(x["cnpj"]) != 14:
            return x["dt"]
        r, f = root_ult.get(x["cnpj"][:8]), fat_root.get(x["cnpj"][:8])
        fd = dt.datetime.strptime(str(f[0]), "%Y%m%d") if f else None
        return max([d for d in (x["dt"], r["dt"] if r else None, fd if fd and fd >= INICIO else None) if d], default=None)

    grupo_cand = C.defaultdict(list)  # cabeca -> [(data, cliente, grupo)]
    for cod, lab in grupos["grupos"].items():
        for x in por_grupo.get(cod, []):
            d = dt_cli(x)
            if d and lab in heads:
                grupo_cand[heads[lab]].append((d, x))
    farmarcas = heads.get("FARMARCAS")
    for lab, h in heads.items():  # cliente do grupo Farmarcas com a bandeira no nome -> cabeca da bandeira
        if lab.endswith(BANDEIRAS_FARMARCAS):
            nuc = sa(lab.replace(BANDEIRAS_FARMARCAS, "")).replace("DROGARIAS ", "").strip()
            for cod in grupos.get("farmarcas", []):
                for x in por_grupo.get(cod, []):
                    t = sa(x["nome"] + " " + x["fant"])
                    if nuc in t or nuc.replace(" ", "") in t.replace(" ", ""):
                        d = dt_cli(x)
                        if d:
                            grupo_cand[h].append((d, x))
    for h, lst in grupo_cand.items():
        d, x = max(lst, key=lambda t: t[0])
        b = best[h]
        if b["dt"] is None or d > b["dt"]:
            b.update(dt=d, fonte="grupo do cadastro", how=f"Grupo do cadastro ({x['gdesc']})", comprador=x["cnpj"],
                     cli=f"{x['nome']} | CNPJ {fmt(x['cnpj'])}")

    # linhas da rede sobem para a cabeca (as bandeiras da Farmarcas depois sobem para a Farmarcas)
    def sobe(i, h):
        b, c = best[h], best[i]
        if c["dt"] and (b["dt"] is None or c["dt"] > b["dt"]):
            if h == farmarcas:
                how = f"Bandeira {s(g(i, 'fant'))}: {c['how']}"
            elif tipo[h].startswith("Associação"):
                how = "Associada da rede"
            else:
                how = "Outra raiz da rede"
            b.update(dt=c["dt"], fonte="linha da rede", how=how, comprador=c["comprador"], cli=c["cli"])
    for i in range(n):
        if i in head_set or not rede[i]:
            continue
        for p in partes(rede[i]):
            sobe(i, heads[p])
    if farmarcas is not None:
        for lab, h in heads.items():
            if lab.endswith(BANDEIRAS_FARMARCAS):
                sobe(h, farmarcas)

    status = ["ATIVO" if b["dt"] and (HOJE - b["dt"]).days <= 360 else "A TRABALHAR" for b in best]

    # ---- 2. vendedor
    real = lambda code: bool(code) and code != "000000" and vname.get(code, "") not in NAO_PESSOA

    def cad(c14):
        if len(c14) != 14:
            return None, None
        if c14 in cad_cnpj and real(cad_cnpj[c14]):
            return vname[cad_cnpj[c14]], "cadastro (CNPJ)"
        for code, _ in cad_root.get(c14[:8], C.Counter()).most_common():
            if real(code):
                return vname[code], "cadastro (raiz)"
        return None, None

    def ult_venda(c14):
        v = fat_root.get(c14[:8]) if len(c14) == 14 else None
        if v and v[1]:
            nomev = max(v[1].items(), key=lambda kv: kv[1])[0]
            if nomev.strip() not in NAO_PESSOA:
                return nomev.strip(), "última venda"
        return None, None
    vend, vend_how = [], C.Counter()
    for i in range(n):
        comp = best[i]["comprador"] if best[i]["dt"] else ""
        v, how = None, None
        for f, c in ((cad, comp), (cad, cnpj[i]), (ult_venda, comp), (ult_venda, cnpj[i])):
            if c:
                v, how = f(c)
                if v:
                    break
        vend.append(v or "")
        vend_how[how] += 1

    # ---- 3. porte
    if vf is None:
        porte = [s(g(i, "porte")) for i in range(n)]
        lojas = [g(i, "lc") if isinstance(g(i, "lc"), (int, float)) else 0 for i in range(n)]
        porte_fonte = C.Counter({"aba": n})
    else:
        porte, lojas, porte_fonte = [], [], C.Counter()
        cab_de = {i: lab for lab, i in heads.items()}
        for i in range(n):
            # cabeca de rede -> a cabeca da mesma rede na VF; senao a mesma raiz com o mesmo tipo; senao rede+tipo+fantasia
            ks = ([("cab", cab_de[i])] if i in cab_de else []) +                  ([("raiz", cnpj[i][:8], tipo[i])] if len(cnpj[i]) == 14 else []) + [("nome", rede[i], tipo[i], sa(g(i, "fant")))]
            k = next((k for k in ks if k in vf), None)
            if k:
                porte.append(vf[k][0]); lojas.append(vf[k][1]); porte_fonte["igual à VF"] += 1
            else:
                lr, lp = g(i, "lr"), g(i, "lp")
                lr = lr if isinstance(lr, (int, float)) else 0
                lp = lp if isinstance(lp, (int, float)) else 0
                lj = max(lr, lp) if tipo[i] != "Farmácia" else lr
                porte.append(porte_por(None, lj)); lojas.append(lj); porte_fonte["regra (lojas)"] += 1

    # ---- 4. caixas, sem contar a mesma loja duas vezes
    receita = [g(i, "lr") if isinstance(g(i, "lr"), (int, float)) else 0 for i in range(n)]
    cobertas = C.Counter()
    for i in range(n):
        if i not in head_set and rede[i]:
            p = partes(rede[i])
            if p:
                cobertas[heads[p[0]]] += lojas[i]
    liq = list(lojas)
    for lab, h in heads.items():
        if h == farmarcas:
            continue
        # a base e o universo da Receita: toda associada ja tem linha propria (com ou sem a bandeira no nome).
        # Associacao/franquia -> so as lojas da propria raiz (CNPJ proprio) ou nenhuma; rede com CNPJ proprio -> nº de
        # lojas considerado (o da rede inteira) menos as outras raizes que ja tem linha
        if tipo[h] == "Associação / franquia (CNPJ próprio)":
            liq[h] = receita[h]
        elif tipo[h].startswith("Associação"):
            liq[h] = 0
        else:
            liq[h] = max(0, lojas[h] - cobertas[h])
    if farmarcas is not None:
        liq[farmarcas] = 0  # rede-mae: as lojas estao nas bandeiras e nas associadas
    cx = [CX.get(porte[i][:1], 0) * (liq[i] or 0) for i in range(n)]
    cx_bruto = [CX.get(porte[i][:1], 0) * (lojas[i] or 0) for i in range(n)]

    return dict(n=n, best=best, status=status, vend=vend, vend_how=vend_how, porte=porte, lojas=lojas, liq=liq, cx=cx,
                cx_bruto=cx_bruto, porte_fonte=porte_fonte, heads=heads, tipo=tipo, rede=rede, cnpj=cnpj,
                antes=[(g(i, "dt"), s(g(i, "st"))) for i in range(n)], homonimos=homonimos)


def grava(nome, R):
    L = LAYOUT[nome]
    cel = {}
    serial = lambda d: (d - dt.datetime(1899, 12, 30)).days
    for i in range(R["n"]):
        r = i + 2
        b = R["best"][i]
        d0, st0 = R["antes"][i]
        if b["fonte"] != "aba":
            cel[(L["col_dt"], r)] = (serial(b["dt"]) if b["dt"] else None, 107)
            cel[(L["col_how"], r)] = (b["how"], None)
            cel[(L["col_cli"], r)] = (b["cli"], None)
        if R["status"][i] != st0:
            cel[(L["col_st"], r)] = (R["status"][i], None)
        nv = L["novas"]
        cel[(nv["vend"], r)] = (R["vend"][i] or None, None)
        if "porte" in nv:
            cel[(nv["porte"], r)] = (R["porte"][i], None)
        cel[(nv["cx"], r)] = (int(R["cx"][i]), 127)
    nv = L["novas"]
    cel[(nv["vend"], 1)] = ("VENDEDOR", L["cab"])
    if "porte" in nv:
        cel[(nv["porte"], 1)] = ("PORTE", L["cab"])
    cel[(nv["cx"], 1)] = ("Total de caixas/ mês", L["cab"])
    larg = {nv["vend"]: 36, nv["cx"]: 14}
    if "porte" in nv:
        larg[nv["porte"]] = 30
    return gravar_celulas(MAPA, nome, cel, larguras=larg)


def main(gravar=False):
    F = fontes()
    grupos = json.load(open(GRUPOS, encoding="utf8"))
    abas = ler_abas()
    vf = processa('"FARMACIA VF"', abas['"FARMACIA VF"'], F, grupos)
    # porte e lojas da VF para a CONSOLIDADO: por raiz; sem CNPJ, por (rede, tipo, fantasia)
    LV = LAYOUT['"FARMACIA VF"']
    rv = abas['"FARMACIA VF"'][1:vf["n"] + 1]
    chave = {}
    for lab, i in vf["heads"].items():
        chave[("cab", lab)] = (vf["porte"][i], vf["lojas"][i])
    for i, r in enumerate(rv):
        if len(vf["cnpj"][i]) == 14:
            chave.setdefault(("raiz", vf["cnpj"][i][:8], vf["tipo"][i]), (vf["porte"][i], vf["lojas"][i]))
        chave.setdefault(("nome", vf["rede"][i], vf["tipo"][i], sa(r[LV["fant"]])), (vf["porte"][i], vf["lojas"][i]))
    co = processa('"FARMACIA CONSOLIDADO"', abas['"FARMACIA CONSOLIDADO"'], F, grupos, chave)
    pickle.dump(dict(vf=vf, co=co), open("farm_final.pkl", "wb"))
    for nome, R in (("CONSOLIDADO", co), ("VF", vf)):
        print("=====", nome, R["n"], "linhas")
        print(" fonte da data:", C.Counter(b["fonte"] for b in R["best"]), " homônimos:", R["homonimos"])
        print(" status antes:", C.Counter(a[1] for a in R["antes"]), " depois:", C.Counter(R["status"]))
        print(" vendedor:", R["vend_how"])
        print(" porte:", C.Counter(p[:1] for p in R["porte"]), R["porte_fonte"])
        print(" caixas (liq):", sum(R["cx"]), " sem o ajuste:", sum(R["cx_bruto"]))
    if gravar:
        try:
            open(MAPA, "r+b").close()
        except PermissionError:
            raise SystemExit("BLOQUEADO - feche o MAPA MERCADO no Excel")
        print(grava('"FARMACIA CONSOLIDADO"', co))
        print(grava('"FARMACIA VF"', vf))


if __name__ == "__main__":
    main("--gravar" in sys.argv)
