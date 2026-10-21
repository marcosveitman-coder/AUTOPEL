"""Grava a base VF (vf.pkl, de vf_build.py) na aba vazia "FARMACIA VF" do MAPA MERCADO, rankeada pelo nº de lojas.
So troca o XML dessa aba (a pasta tem tabelas dinamicas: nao salvar pelo openpyxl).  Rodar na pasta onde esta o vf.pkl."""
import os, re, sys, pickle, zipfile
from openpyxl.utils import get_column_letter
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from grava_consolidado import fmt, excel_date, cell, S_HEAD, S_TXT, S_NUM, S_DATE

MAPA = os.environ.get("AUTOPEL_MAPA", "MAPA MERCADO.xlsx")
NOME_ABA = "&quot;FARMACIA VF&quot;"
COLS = [("POSIÇÃO", 8), ("REDE / ASSOCIAÇÃO", 34), ("TIPO", 24), ("ENTIDADE / FEDERAÇÃO", 16), ("NOME FANTASIA", 40),
        ("RAZÃO SOCIAL", 48), ("CNPJ", 20), ("Nº DE LOJAS (CONSIDERADO)", 12), ("Nº DE LOJAS (RECEITA)", 12),
        ("Nº DE LOJAS (PESQUISA / RANKING)", 12), ("UF SEDE", 7), ("MUNICÍPIO SEDE", 22), ("SITUAÇÃO NA RECEITA", 12),
        ("VALIDAÇÃO / OBSERVAÇÃO", 60), ("DATA ÚLTIMA COMPRA", 13), ("STATUS", 13), ("COMO FOI ACHADO NO FATURAMENTO", 28),
        ("CLIENTE ENCONTRADO NO FATURAMENTO", 60), ("FATURAMENTO DA FARMÁCIA 2025 (R$, RANKING)", 18),
        ("VALOR FATURADO AUTOPEL 2025 (R$)", 16), ("PORTE (RECEITA)", 22)]


def sheet_path(z):
    wb = z.read("xl/workbook.xml").decode(); rels = z.read("xl/_rels/workbook.xml.rels").decode()
    rid = re.search(r'<sheet [^>]*name="%s"[^>]*r:id="([^"]*)"' % re.escape(NOME_ABA), wb).group(1)
    t = re.search(r'<Relationship [^>]*Id="%s"[^>]*/>' % rid, rels).group(0)
    return "xl/" + re.search(r'Target="/?(?:xl/)?([^"]*)"', t).group(1)


def main():
    rows = pickle.load(open("vf.pkl", "rb"))
    rows.sort(key=lambda x: (-(x["lojas_cons"] or 0), -(x["fat_rede"] or 0), x["fant"] or ""))
    out = [f'<row r="1" spans="1:{len(COLS)}" ht="45" customHeight="1">' +
           "".join(cell(f"{get_column_letter(j+1)}1", h, S_HEAD) for j, (h, w) in enumerate(COLS)) + "</row>"]
    n = 1
    for x in rows:
        n += 1
        u, rk = x.get("ult"), x.get("rk") or {}
        cli = ""
        if u:
            cli = f'{u.get("cliente", "")} | CNPJ {fmt(u.get("cnpj", ""))}'
            if u.get("associada"): cli += f' | associada: {u["associada"]}'
            if u.get("produto"): cli += f' | produto: {u["produto"]}'
            if u.get("dominio"): cli += f' | e-mail @{u["dominio"]}'
        num = lambda v: round(v, 2) if isinstance(v, (int, float)) and v else None
        vals = [(n - 1, S_NUM), (x["rede"], S_TXT), (x["tipo"], S_TXT), (rk.get("entidade"), S_TXT), (x["fant"], S_TXT),
                (x["rz"], S_TXT), (fmt(x["cnpj"]), S_TXT), (x["lojas_cons"], S_NUM),
                (x["lojas"] if x["lojas"] is not None else None, S_NUM),
                (rk.get("lojas") if isinstance(rk.get("lojas"), (int, float)) else (x["inf"] if isinstance(x["inf"], (int, float)) else None), S_NUM),
                (x["uf"], S_TXT), (x["mun"], S_TXT), (x["sit"], S_TXT), (x["obs"], S_TXT),
                (excel_date(u["dt"]) if u else None, S_DATE), (x["status"], S_TXT), (u["how"] if u else "", S_TXT), (cli, S_TXT),
                (num(x["fat_rede"]), S_NUM), (num(x["fat25"]), S_NUM), (x["porte"], S_TXT)]
        out.append(f'<row r="{n}" spans="1:{len(COLS)}">' + "".join(cell(f"{get_column_letter(j+1)}{n}", v, s) for j, (v, s) in enumerate(vals)) + "</row>")
    last = get_column_letter(len(COLS))
    cols = "<cols>" + "".join(f'<col min="{j+1}" max="{j+1}" width="{w}" customWidth="1"/>' for j, (h, w) in enumerate(COLS)) + "</cols>"
    z = zipfile.ZipFile(MAPA)
    SHEET = sheet_path(z)
    x = z.read(SHEET).decode("utf8")
    ja = re.search(r"<cols>.*?</cols><sheetData>.*?</sheetData><autoFilter[^>]*/>", x, re.S)
    if ja and "POSI" in x[:20000]:  # aba ja gravada por este script: refaz
        x = x.replace(ja.group(0), "<sheetData/>")
        x = re.sub(r'<dimension ref="[^"]*"/>', '<dimension ref="A1"/>', x)
        x = re.sub(r"<sheetViews>.*?</sheetViews>", '<sheetViews><sheetView workbookViewId="0"/></sheetViews>', x, flags=re.S)
    assert "<sheetData/>" in x, "a aba FARMACIA VF tem outro conteudo - nao sobrescrevo"
    x = x.replace('<dimension ref="A1"/>', f'<dimension ref="A1:{last}{n}"/>')
    x = re.sub(r"<sheetViews>.*?</sheetViews>",
               '<sheetViews><sheetView workbookViewId="0"><pane ySplit="1" topLeftCell="A2" activePane="bottomLeft" state="frozen"/>'
               '<selection pane="bottomLeft" activeCell="A2" sqref="A2"/></sheetView></sheetViews>', x, count=1, flags=re.S)
    x = x.replace("<sheetData/>", cols + "<sheetData>" + "".join(out) + "</sheetData>" + f'<autoFilter ref="A1:{last}{n}"/>')
    tmp = MAPA + ".tmp"
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zo:
        for info in z.infolist():
            zo.writestr(info, x.encode("utf8") if info.filename == SHEET else z.read(info.filename), compress_type=zipfile.ZIP_DEFLATED)
    z.close()
    os.replace(tmp, MAPA)
    print("gravado:", n - 1, "linhas em", SHEET)


if __name__ == "__main__":
    try:
        open(MAPA, "r+b").close()
    except PermissionError:
        raise SystemExit("BLOQUEADO - feche o MAPA MERCADO no Excel")
    main()
