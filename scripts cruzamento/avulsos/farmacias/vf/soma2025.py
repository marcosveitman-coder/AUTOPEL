"""Soma o VLR_TOTAL faturado em 2025 por raiz de CNPJ, por COD_CLI e por (DES_PROD -> raiz).  Saida: soma2025.pkl"""
import os, sys, glob, re, pickle, time, collections
from multiprocessing import Pool
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", ".."))
from common import BASE
from extract import to_int_date
NEED = ["EMISSAO", "COD_CLI", "CNPJ", "DES_PROD", "VLR_TOTAL"]

def work(job):
    f, sh = job
    import openpyxl
    wb = openpyxl.load_workbook(os.path.join(BASE, f), read_only=True)
    it = wb[sh].iter_rows(values_only=True)
    hdr = [str(h).strip() if h is not None else "" for h in next(it, ())]
    if any(n not in hdr for n in NEED):
        return f, {}, {}, {}
    iE, iK, iC, iP, iT = (hdr.index(n) for n in NEED)
    root, cod, desc = collections.Counter(), collections.Counter(), collections.defaultdict(collections.Counter)
    for r in it:
        dt = to_int_date(r[iE])
        if not dt or dt // 10000 != 2025:
            continue
        try: v = float(r[iT] or 0)
        except (TypeError, ValueError): continue
        c = re.sub(r"\D", "", str(r[iC] or ""))
        if len(c) == 14: root[c[:8]] += v
        k = str(r[iK] or "").strip()
        if k: cod[k.zfill(6)] += v
        p = str(r[iP] or "").strip()
        if p and len(c) == 14: desc[p][c[:8]] += v
    return f + " " + sh, root, cod, desc

if __name__ == "__main__":
    import openpyxl
    jobs = []
    for f in sorted(os.path.basename(p) for p in glob.glob(os.path.join(BASE, "fat 202*.xlsx"))):
        if not re.search(r"fat 202[56]", f): continue
        wb = openpyxl.load_workbook(os.path.join(BASE, f), read_only=True); jobs += [(f, s) for s in wb.sheetnames]; wb.close()
    t = time.time(); R, C, Dd = collections.Counter(), collections.Counter(), collections.defaultdict(collections.Counter)
    with Pool(4) as p:
        for name, r, c, d in p.imap_unordered(work, jobs):
            R.update(r); C.update(c)
            for k, v in d.items(): Dd[k].update(v)
            print(name, "raizes 2025:", len(r), "valor:", round(sum(r.values())), f"{time.time()-t:.0f}s", flush=True)
    pickle.dump({"root": dict(R), "cod": dict(C), "desc": {k: dict(v) for k, v in Dd.items()}}, open("soma2025.pkl", "wb"))
    print("TOTAL raizes:", len(R), "valor 2025:", round(sum(R.values())), f"{time.time()-t:.0f}s")
