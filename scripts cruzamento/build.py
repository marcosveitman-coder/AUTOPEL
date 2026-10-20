"""Escolhe o resultado final por supermercado e grava a planilha nova com 4 colunas no fim."""
import os, pickle, datetime, collections
from copy import copy
import openpyxl
from openpyxl.utils import get_column_letter
from common import SM_PATH, OUT_PATH, CAND, pick_vendor, fmt_cnpj

LABEL = {"cnpj": "CNPJ raiz", "razao": "Razão social", "fantasia": "Nome fantasia + UF", "desc": "Descrição do produto"}
ORDER = ["cnpj", "razao", "fantasia", "desc"]


def to_date(dt):
    return datetime.datetime(dt // 10000, dt // 100 % 100, dt % 100)


def choose(c):
    """CNPJ/razao social (fortes) tem prioridade; nome fantasia/descricao (fracos) so entram se nao houver forte,
    ou se forem mais recentes E o proprio cliente comprador tiver o nome da rede."""
    strong = [(m, c[m]) for m in ("cnpj", "razao") if m in c]
    weak = [(m, c[m]) for m in ("fantasia", "desc") if m in c]
    pick = lambda lst: max(lst, key=lambda md: (md[1]["dt"], -ORDER.index(md[0])))
    if not strong:
        return pick(weak) if weak else None
    best = pick(strong)
    over = [md for md in weak if md[1]["buyer_has_brand"] and md[1]["dt"] > best[1]["dt"]]
    return pick(over) if over else best


def describe(m, d):
    s = f'{d["cliente"]} | CNPJ {fmt_cnpj(d["cnpj"])}'
    if d.get("seg"):
        s += f' | {d["seg"]}'
    if m == "desc":
        s += f' | Produto: {d["produto"]}'
    return s


def main():
    with open(CAND, "rb") as f:
        cand = pickle.load(f)["cand"]
    wb = openpyxl.load_workbook(SM_PATH)
    ws = wb.active
    assert ws.max_row - 1 == len(cand), (ws.max_row, len(cand))
    first = ws.max_column + 1  # 27 = AA
    heads = ["Última venda (faturamento)", "Vendedor da última venda", "Critério do cruzamento", "Cliente encontrado no faturamento"]
    widths = [16, 40, 22, 90]
    src = ws.cell(1, ws.max_column)
    for k, h in enumerate(heads):
        c = ws.cell(1, first + k, h)
        c.font, c.fill, c.border, c.alignment = copy(src.font), copy(src.fill), copy(src.border), copy(src.alignment)
        ws.column_dimensions[get_column_letter(first + k)].width = widths[k]
    stats = collections.Counter()
    for i, cd in enumerate(cand):
        r = i + 2
        ch = choose(cd)
        if not ch:
            stats["sem venda"] += 1
            continue
        m, d = ch
        stats[LABEL[m]] += 1
        dc = ws.cell(r, first, to_date(d["dt"]))
        dc.number_format = "DD/MM/YYYY"
        ws.cell(r, first + 1, pick_vendor(d["vends"]) or "(sem vendedor no faturamento)")
        ws.cell(r, first + 2, LABEL[m])
        ws.cell(r, first + 3, describe(m, d))
    ws.auto_filter.ref = f"A1:{get_column_letter(first + 3)}{ws.max_row}"
    try:
        wb.save(OUT_PATH)
    except PermissionError:
        raise SystemExit(f"Nao consegui salvar {OUT_PATH} - feche o arquivo no Excel e rode de novo.")
    print("salvo:", OUT_PATH)
    for k, v in stats.most_common():
        print(f"  {k}: {v}")


if __name__ == "__main__":
    main()
